"""GridMats Solar Walkway proposal template: cost-benefit workbook.

Usage: python3 build_sw_model.py ../clients/<id>.json <out.xlsx>
Every client-specific value lives on the 'Client' sheet; GridMats product and cost
assumptions live on 'Product & Costs'. All results are formulas.
"""
import json
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import xlhelpers as X
from xlhelpers import (BOX, FMT_INR, FMT_INR2, FMT_NUM, FMT_NUM1, FMT_NUM2, FMT_PCT, FMT_X, FMT_YRS, FONT, NOTE_FONT, TOP2,
                       header, put, section, title, widths)

# Solar Walkway palette (GridMats brick red + dark slate)
X.GREEN = "1F2A33"
X.GREEN_LIGHT = "FBEAEC"
X.GREEN_DARK = "1A1A1A"

CFG = json.load(open(sys.argv[1]))
OUT = sys.argv[2]
wb = openpyxl.Workbook()
N = {}
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
DAYS = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
FMT_LAKH = '₹#,##0.00" L";(₹#,##0.00" L");"-"'


def dname(name, sheet, ref):
    s = f"'{sheet}'" if any(ch in sheet for ch in " -&") else sheet
    col = "".join(ch for ch in ref if ch.isalpha())
    row = "".join(ch for ch in ref if ch.isdigit())
    wb.defined_names[name] = DefinedName(name, attr_text=f"{s}!${col}${row}")
    N[name] = (sheet, f"{col}{row}")


def note(ws, ref, text):
    put(ws, ref, text, kind="note")
    ws[ref].alignment = Alignment(wrap_text=True, vertical="top")


cover = wb.active
cover.title = "Cover"
cl = wb.create_sheet("Client")
pc = wb.create_sheet("Product & Costs")
cb = wb.create_sheet("Cost Build-up")
mt = wb.create_sheet("Module Tech")
irr = wb.create_sheet("Irradiance")
en = wb.create_sheet("Energy")
boq = wb.create_sheet("Pilot BoQ")
pe = wb.create_sheet("Pilot Economics")
ro = wb.create_sheet("Rollout Economics")
cmpw = wb.create_sheet("Comparison")
sens = wb.create_sheet("Sensitivity")
sc = wb.create_sheet("Scale-up")
card = wb.create_sheet("Scorecard")
src = wb.create_sheet("Sources & Notes")
for w in wb.worksheets:
    w.sheet_properties.tabColor = {"Cover": "C22935", "Client": "FFC000", "Product & Costs": "FFC000", "Module Tech": "FFC000", "Pilot BoQ": "C22935", "Cost Build-up": "7F7F7F"}.get(w.title, "1F2A33")


# ---------------------------------------------------------------- generic input-table writer
class InputTable:
    def __init__(self, ws, sheet, start_row, scenario=False):
        self.ws, self.sheet, self.row, self.scn = ws, sheet, start_row, scenario

    def sec(self, text, ncols=7):
        self.row += 1
        section(self.ws, self.row, text, ncols=ncols)
        self.row += 1

    def add(self, name, label, unit, val, fmt, source, key=False, dv=None, link=False):
        ws, r = self.ws, self.row
        put(ws, f"A{r}", label, kind="label")
        put(ws, f"B{r}", unit, kind="note")
        if isinstance(val, tuple):
            for col, v in zip("DEF", val):
                put(ws, f"{col}{r}", v, fmt, kind="input", key=key)
            put(ws, f"C{r}", f"=INDEX(D{r}:F{r},ScnIdx)", fmt, bold=True)
        elif isinstance(val, str) and val.startswith("="):
            put(ws, f"C{r}", val, fmt, kind="calc")
        else:
            put(ws, f"C{r}", val, fmt, kind="input", key=key, align="left" if isinstance(val, str) else None)
            if dv:
                d = DataValidation(type="list", formula1=dv, allow_blank=False)
                ws.add_data_validation(d)
                d.add(f"C{r}")
        note(ws, f"G{r}", source)
        dname(name, self.sheet, f"C{r}")
        self.row += 1


# ================================================================ CLIENT
ws = cl
title(ws, f"Client and site: {CFG['client_short']}",
      "Everything that changes from one client to the next is on this sheet. Yellow cells are the ones to confirm first.")
widths(ws, {"A": 50, "B": 14, "C": 30, "D": 13, "E": 13, "F": 13, "G": 80})
header(ws, 4, ["Item", "Unit", "Value used", "Conservative", "Base", "Optimistic", "Source / note"])
put(ws, "A5", "Scenario used across the model", kind="label", bold=True)
put(ws, "C5", "Base", kind="input", key=True, bold=True, align="center")
d = DataValidation(type="list", formula1='"Conservative,Base,Optimistic"', allow_blank=False)
ws.add_data_validation(d); d.add("C5")
put(ws, "D5", "Conservative", kind="note"); put(ws, "E5", "Base", kind="note"); put(ws, "F5", "Optimistic", kind="note")
note(ws, "G5", "Switches performance ratio and degradation on 'Product & Costs'.")
dname("Scenario", "Client", "C5")
put(ws, "A6", "Scenario index", kind="label")
put(ws, "C6", "=MATCH(C5,D5:F5,0)", FMT_NUM)
dname("ScnIdx", "Client", "C6")
t = InputTable(ws, "Client", 7)
t.sec("1. Client and site")
t.add("ClientName", "Client (short name used on slides)", "text", CFG["client_short"], None, "Shown on the cover and slides.", key=True)
t.add("ClientEntity", "Client legal entity", "text", CFG["client_entity"], None, "Use the entity that will issue the PO.", key=True)
t.add("SiteName", "Campus / site", "text", CFG["site_name"], None, "Confirm at the site walk.", key=True)
t.add("SiteCity", "City, state", "text", CFG["site_city"], None, "")
t.add("IrrCity", "Sunshine profile (city)", "choice", CFG["irradiance_city"], None,
      "Monthly sunshine from the 'Irradiance' sheet. Replace with site data after the site assessment.", key=True, dv="=Irradiance!$B$7:$B$15")
t.add("Utility", "Electricity supplier", "text", CFG["utility"], None, "")
t.add("TariffCat", "Tariff category", "text", CFG["tariff_category"], None, "Confirm from the campus electricity bill.")
t.sec("2. Power price and site losses")
t.add("Tariff", "Grid power cost avoided, year 1 (landed)", "₹ / kWh", CFG["tariff"], FMT_INR2, CFG["tariff_note"], key=True)
t.add("TariffEsc", "Power tariff escalation", "% / yr", CFG["tariff_escalation"], FMT_PCT, "Assumption.")
t.add("ShadeLoss", "Site shading loss (trees, buildings)", "%", CFG["shading_loss"], FMT_PCT, "Assumption until the site assessment measures it.")
t.sec("3. Pilot and rollout size")
t.add("BuildType", "Build type", "choice", CFG["build_type"], None,
      "'New walkway' builds the base and credits the paver walkway it replaces; 'Overlay on existing paving' levels the existing path and adds edge ramps.",
      dv='"New walkway,Overlay on existing paving"')
t.add("SoilCBR", "Soil strength under the walkway: soaked CBR", "%", CFG["soil_cbr"], FMT_NUM1, CFG["soil_cbr_note"], key=True)
t.add("PilotL", "Pilot walkway length", "m", CFG["pilot_length_m"], FMT_NUM1, "Pilots are 20-25 m².", key=True)
t.add("PilotW", "Pilot walkway width", "m", CFG["pilot_width_m"], FMT_NUM1, "", key=True)
t.add("PilotArea", "Pilot area", "m²", "=PilotL*PilotW", FMT_NUM1, "Length × width.")
t.add("Phase2Area", "Phase 2 area (after the pilot)", "m²", CFG["phase2_area"], FMT_NUM, "Illustrative.")
t.add("CampusArea", "Campus rollout area", "m²", CFG["campus_area"], FMT_NUM, "Used on 'Rollout Economics' and 'Comparison'.", key=True)
t.add("PortfolioArea", "Portfolio area", "m²", CFG["portfolio_area"], FMT_NUM, CFG["portfolio_label"])
t.sec("4. Commercial and finance")
t.add("CoFund", "GridMats co-funding of the pilot PV layer", "%", CFG["cofund"], FMT_PCT, "Pilot offer. Needs MD approval per pilot.", key=True)
t.add("DiscRate", "Discount rate", "%", CFG["discount_rate"], FMT_PCT, "Client hurdle rate; 10% if not given.")
t.add("TaxRate", "Client income-tax rate", "%", CFG["tax_rate"], FMT_PCT, "Section 115BAA rate; used for the after-tax view.")
t.add("EvalMonths", "Evaluation period before the rollout decision", "months", CFG["evaluation_months"], FMT_NUM, "Earlier Jio proposal letter: 3-month pilot evaluation.")
t.add("MonitorMonths", "Monitoring period", "months", CFG["monitoring_months"], FMT_NUM, "Covers one monsoon.")
t.add("PropDate", "Proposal date", "text", CFG["proposal_date"], None, "")
t.add("PreparedBy", "Prepared by", "text", CFG["prepared_by"], None, "")
ws.freeze_panes = "A5"

# ================================================================ PRODUCT & COSTS
ws = pc
title(ws, "GridMats Solar Walkway: product and cost assumptions",
      "Standard GridMats inputs, the same for every client. ₹ shown per sq ft first, then per m², ex-GST.")
widths(ws, {"A": 60, "B": 15, "C": 26, "D": 13, "E": 13, "F": 13, "G": 90})
header(ws, 4, ["Assumption", "Unit", "Value used", "Conservative", "Base", "Optimistic", "Source / note"])
t = InputTable(ws, "Product & Costs", 4)
t.sec("1. Product: module technology and rated power")
t.add("TechSel", "Cell technology of the walkable module", "choice", "TOPCon (n-type)", None,
      "Pick from the 'Module Tech' library. Drives temperature loss and degradation; confirm with the module supplier's datasheet.", key=True,
      dv="='Module Tech'!$B$6:$B$8")
t.add("PowerBasis", "Rated-power basis", "choice", "Design basis (40 W per module)", None,
      "Design basis = the MD-approved 40 W per module. 'Technology-derived' rates the module from its cell technology (upgrade path, to be validated).",
      dv='"Design basis (40 W per module),Technology-derived"')
t.add("ModWBasis", "Design-basis module rated power", "Wp", 40, FMT_NUM, "GridMats Solar Walkway flier (6 Oct 2026) and Client/Partner Deck v2.")
t.add("ModPerM2", "Modules per m² of walkway (one on each 500 × 500 mm grid)", "nos / m²", 4, FMT_NUM, "Flier (6 Oct 2026): 4 grids per m².")
t.add("ModL", "Module length", "mm", 550, FMT_NUM, "Client/Partner Deck v2.")
t.add("ModWd", "Module width", "mm", 400, FMT_NUM, "Client/Partner Deck v2.")
t.add("ModArea", "Module area", "m²", "=ModL*ModWd/1000000", FMT_NUM2, "Length × width.")
t.add("TechCellEff", "Cell efficiency, mass production (selected technology)", "%",
      "=INDEX('Module Tech'!$C$6:$C$8,MATCH(TechSel,'Module Tech'!$B$6:$B$8,0))", FMT_PCT, "From 'Module Tech'.")
t.add("TechModEff", "Standard module efficiency (selected technology)", "%",
      "=INDEX('Module Tech'!$E$6:$E$8,MATCH(TechSel,'Module Tech'!$B$6:$B$8,0))", FMT_PCT, "From 'Module Tech'.")
t.add("TechTempCoef", "Temperature coefficient of power", "% / °C",
      "=INDEX('Module Tech'!$F$6:$F$8,MATCH(TechSel,'Module Tech'!$B$6:$B$8,0))", '0.00%', "From 'Module Tech'.")
t.add("TechDeg1", "First-year degradation", "%", "=INDEX('Module Tech'!$G$6:$G$8,MATCH(TechSel,'Module Tech'!$B$6:$B$8,0))", FMT_PCT, "From 'Module Tech'.")
t.add("TechDegAnn", "Annual degradation after year 1", "% / yr", "=INDEX('Module Tech'!$H$6:$H$8,MATCH(TechSel,'Module Tech'!$B$6:$B$8,0))", FMT_PCT, "From 'Module Tech'.")
t.add("WalkPackLoss", "Walk-on build loss vs a standard module (thick textured glass, lower cell packing)", "%", 0.15, FMT_PCT,
      "Assumption for the technology-derived rating. The 40 W design basis implies about 23% on a 23.5% TOPCon module.")
t.add("ModWTech", "Technology-derived module rating", "Wp", "=ModArea*TechModEff*(1-WalkPackLoss)*1000", FMT_NUM1, "Module area × efficiency × (1 − walk-on loss).")
t.add("ModW", "Module rated power used", "Wp", "=IF(PowerBasis=\"Technology-derived\",ModWTech,ModWBasis)", FMT_NUM1, "Per the rated-power basis above.")
t.add("PowerDens", "Rated power of the walkway", "Wp / m²", "=ModW*ModPerM2", FMT_NUM, "About 160 W/m² on the design basis (flier, 6 Oct 2026).")
t.add("ModEff", "Walkway module efficiency (check)", "%", "=ModW/(ModArea*1000)", FMT_PCT, "Wp ÷ (module area × 1,000 W/m²).")
t.add("SurfEff", "Surface efficiency of the walkway", "%", "=PowerDens/1000", FMT_PCT, "Flier: 15-16% surface efficiency.")
t.sec("2. Performance: loss stack for flat, walk-on modules")
t.add("Tcell", "Irradiance-weighted cell temperature on the walkway", "°C", (60, 52, 45), FMT_NUM,
      "Assumption. Pavement-integrated PV runs about 20 °C hotter than open-rack; GridMats' 50 mm gaps ventilate the modules.", key=True)
t.add("LossTemp", "Temperature loss", "%", "=-TechTempCoef*(Tcell-25)", FMT_PCT, "−temperature coefficient × (cell temperature − 25 °C).")
t.add("LossSoil", "Soiling (flat surface, foot traffic; monthly cleaning)", "%", (0.15, 0.08, 0.04), FMT_PCT,
      "Horizontal modules soil about twice as fast as tilted ones; annual losses of 5-18% reported without frequent cleaning.", key=True)
t.add("LossFoot", "Shading by people walking on it", "%", (0.08, 0.04, 0.02), FMT_PCT, "Assumption; depends on footfall at the chosen walkway.")
t.add("LossRefl", "Reflection and angle losses (flat, textured glass)", "%", (0.05, 0.03, 0.02), FMT_PCT, "Assumption.")
t.add("LossMism", "Mismatch and DC wiring", "%", (0.04, 0.03, 0.02), FMT_PCT, "Assumption.")
t.add("LossInv", "Inverter", "%", (0.04, 0.035, 0.025), FMT_PCT, "Assumption.")
t.add("LossAvail", "Availability", "%", (0.03, 0.01, 0.005), FMT_PCT, "Assumption.")
t.add("PR", "Performance ratio (product of the losses above)", "%",
      "=(1-LossTemp)*(1-LossSoil)*(1-LossFoot)*(1-LossRefl)*(1-LossMism)*(1-LossInv)*(1-LossAvail)", FMT_PCT,
      "Earlier ROI model used a flat 78% (tilted C&I). Field results for road- and walk-integrated PV have been about half of rooftop yield per m²; the conservative case reflects that.")
t.add("WearAdder", "Extra degradation from walk-on use", "% / yr", (0.004, 0.002, 0.001), FMT_PCT, "Assumption.")
t.add("Deg1", "First-year degradation used", "%", "=TechDeg1", FMT_PCT, "Selected technology.")
t.add("Degr", "Annual degradation used (technology + walk-on wear)", "% / yr", "=TechDegAnn+WearAdder", FMT_PCT, "Selected technology plus wear.")
t.add("Life", "Analysis horizon and module life", "years", 20, FMT_NUM, "Earlier ROI model: 20 years. Tables are built for 20 years.")
t.add("InvRepl", "Inverter replacement", "₹ / kWp", 4000, FMT_INR, "Assumption.")
t.add("InvYear", "Inverter replacement year", "year", 12, FMT_NUM, "Assumption.")
t.add("OMPct", "Annual O&M (cleaning, inspection, minor repairs)", "% of capex / yr", 0.015, FMT_PCT, "Earlier ROI model.")
t.sec("3. Other prices and costs (ex-GST). PV-layer, electrical and labour prices are built up from actual unit costs on 'Cost Build-up'")
t.add("ScaleFrom", "Rollout prices (volume cost-down, project management) apply from", "m²", 500, FMT_NUM, "Assumption.")
t.add("GridT1", "GridMats OC-LW 40 base grid, under 1,000 m²", "₹ / m²", 690, FMT_INR, "GridMats published price ladder (25 Sep 2026): ₹64/sq ft. Cost is ₹125 per grid (₹500/m²): see 'Cost Build-up'.")
t.add("GridT2", "GridMats OC-LW 40, 1,000-2,500 m²", "₹ / m²", 650, FMT_INR, "₹60/sq ft.")
t.add("GridT3", "GridMats OC-LW 40, 2,500-5,000 m²", "₹ / m²", 620, FMT_INR, "₹58/sq ft.")
t.add("GridT4", "GridMats OC-LW 40, 5,000 m² and above", "₹ / m²", 590, FMT_INR, "₹55/sq ft.")
t.add("CBRLimit", "Soaked CBR below which the thicker WMM layer is used", "%", 8, FMT_NUM1, "GridMats base specification (7 Oct 2026): CBR under 8 → 100 mm WMM; 8 or more → 50 mm WMM, on well-compacted soil.")
t.add("WMMThick", "Wet mix macadam (WMM) thickness used", "mm", "=IF(SoilCBR<CBRLimit,100,50)", FMT_NUM, "From the site's soaked CBR on 'Client'.")
t.add("RateSubgrade", "Subgrade preparation: level, proof-roll and compact to 95-97% MDD", "₹ / m²", 60, FMT_INR, "Design Engine 73 unit rate.")
t.add("RateGeo", "Geotextile fabric (non-woven, laid on the compacted subgrade)", "₹ / m²", 55, FMT_INR, "ASSUMPTION: supply and lay of 200 gsm non-woven geotextile, with overlaps. Confirm the local rate.", key=True)
t.add("RateWMM", "Wet mix macadam, supplied, laid and rolled", "₹ / m³", 2000, FMT_INR, "ASSUMPTION: replaces the ₹1,600/m³ GSB rate in Design Engine 73. Confirm the local rate.", key=True)
t.add("MsandThick", "M-sand bedding thickness", "mm", 20, FMT_NUM, "Design Engine 73 (20 mm bedding).")
t.add("RateMsand", "M-sand bedding, supplied and screeded", "₹ / m³", 1800, FMT_INR, "Design Engine 73 unit rate (₹36/m² at 20 mm).")
t.add("RateEdge", "Edge restraint", "₹ / m²", 12, FMT_INR, "Design Engine 73 unit rate.")
t.add("CivilNew", "New-walkway base works: subgrade, geotextile, WMM, M-sand bedding, edges", "₹ / m²",
      "=RateSubgrade+RateGeo+WMMThick/1000*RateWMM+MsandThick/1000*RateMsand+RateEdge", FMT_INR,
      "Sum of the lines above: ₹363/m² on soft soil (100 mm WMM); ₹263/m² otherwise (50 mm WMM) at the rates shown.")
t.add("CivilOverlay", "Overlay on existing paving (levelling, edge ramps)", "₹ / m²", 150, FMT_INR, "Assumption.")
t.add("IntegLump", "Grid integration and metering, per connection point", "₹", 25000, FMT_INR, "Earlier ROI model (pilot scale).")
t.add("ConnKWp", "kWp per connection point at rollout", "kWp", 50, FMT_NUM, "Assumption.")
t.add("LogPct", "Logistics and freight", "% of hardware", 0.05, FMT_PCT, "Earlier ROI model.")
t.add("ConvWalk", "Paver / concrete walkway it replaces (new build only)", "₹ / m²", 1200, FMT_INR, "Earlier ROI model: ₹800-1,500/m² range.")
t.sec("4. Pilot services (valued; offered free in the pilot)")
t.add("DesignLump", "Site assessment, shading study and engineering", "₹", 40000, FMT_INR, "Estimate of value.")
t.add("SupDays", "GridMats supervision and training on site", "days", 3, FMT_NUM, "Standard: up to 3 days.")
t.add("ManDay", "Supervisor man-day rate", "₹ / day", 15000, FMT_INR, "GridMats standard rate.")
t.add("TravelLump", "Supervisor travel and stay", "₹", 40000, FMT_INR, "Estimate; normally paid by the client.")
t.add("IoTPkg", "IoT sensors, live dashboard and performance study for the monitoring period", "₹", 125000, FMT_INR,
      "Estimate of value: module-temperature and moisture sensors, reference cell, energy meter and logger, dashboard, reports.")
t.sec("5. Comparison benchmarks")
t.add("RoofWp", "C&I rooftop solar, all-in", "₹ / Wp", 40, FMT_INR2, "Earlier ROI model cost stack (₹35-45/Wp).")
t.add("RoofDens", "Rooftop / carport panel output per m² of panel", "Wp / m²", 190, FMT_NUM, "19% efficient panels (earlier ROI model).")
t.add("RoofPR", "Rooftop / carport performance ratio", "%", 0.78, FMT_PCT, "Earlier ROI model.")
t.add("RoofDegr", "Rooftop / carport annual degradation", "% / yr", "=TechDegAnn", FMT_PCT, "Same cell technology as the walkway, without walk-on wear.")
t.add("RoofOM", "Rooftop O&M", "% of capex / yr", 0.01, FMT_PCT, "Earlier ROI model.")
t.add("CarportWp", "Solar carport / canopy, all-in", "₹ / Wp", 61, FMT_INR2, "Earlier ROI model (₹55-65/Wp).")
t.add("CarportOM", "Carport O&M", "% of capex / yr", 0.012, FMT_PCT, "Earlier ROI model.")
t.add("ConvWalkOM", "Conventional walkway maintenance", "% of capex / yr", 0.02, FMT_PCT, "Earlier ROI model.")
t.sec("6. Tax and environment")
t.add("GSTRate", "GST on a solar walkway system (composite)", "%", 0.089, FMT_PCT,
      "5% on solar goods from 22 Sep 2025 (Notification 9/2025-CT(R)) on 70% of value + 18% on 30% services = 8.9%. Earlier model used 13.8%. Input tax credit on captive solar allowed per Gujarat AAR (Jan 2024); confirm for the client.")
t.add("ADRate", "Accelerated depreciation on solar (WDV)", "%", 0.40, FMT_PCT, "Income-tax Act: 40% WDV.")
t.add("GridEF", "India grid emission factor", "t CO2 / MWh", 0.710, FMT_NUM2, "CEA CO2 Baseline Database v21.0 (FY2024-25).")
t.add("PlasticLW", "Recycled PP in the OC-LW 40 base", "kg / m²", 4.0, FMT_NUM1, "GridMats master data.")
ws.freeze_panes = "A5"


# ================================================================ COST BUILD-UP (GridMats internal)
ws = cb
title(ws, "Cost build-up: GridMats internal. Remove this sheet before sending the workbook to a client",
      "Starts from GridMats' actual unit costs per 40 W piece, adds what those costs leave out, and builds the client price. Margin is shown on this sheet only.")
widths(ws, {"A": 62, "B": 15, "C": 26, "D": 15, "E": 15, "F": 15, "G": 90})
header(ws, 4, ["Item", "Unit", "Value used", "Conservative", "Base", "Optimistic", "Source / note"])
t = InputTable(ws, "Cost Build-up", 4)
t.sec("1. GridMats actual material cost per 40 W piece (one piece on each 500 × 500 mm grid)")
t.add("CostModule", "Solar module, 40 W (550 × 400 mm)", "₹ / piece", 1200, FMT_INR,
      "GridMats actual cost (7 Oct 2026). Excludes margin, labour, wiring, inverter and management.", key=True)
t.add("CostGlass", "12 mm toughened glass", "₹ / piece", 1600, FMT_INR, "GridMats actual cost (7 Oct 2026).", key=True)
t.add("CostGrid", "GridMats OC-LW 40 base grid", "₹ / piece", 125, FMT_INR, "GridMats actual cost (7 Oct 2026). Sold at the published ladder price, not at cost-plus.", key=True)
t.add("CostLumber", "MLP lumber (40 × 40 mm recycled)", "₹ / piece", 150, FMT_INR, "GridMats actual cost (7 Oct 2026).", key=True)
t.add("MatPiece", "Material cost per piece", "₹ / piece", "=CostModule+CostGlass+CostGrid+CostLumber", FMT_INR, "Sum of the four lines above.")
t.add("MatM2", "Material cost per m² of walkway", "₹ / m²", "=MatPiece*ModPerM2", FMT_INR, "Pieces per m² × material cost per piece.")
t.add("MatSqft", "Material cost per sq ft", "₹ / sq ft", "=MatM2/10.7639", FMT_INR, "")
t.add("MatWp", "Material cost per Wp", "₹ / Wp", "=MatPiece/ModW", FMT_INR2, "Rooftop solar is about ₹40/Wp all-in, so the walkway's material cost alone is about twice that.")
t.sec("2. Not in the material cost. ASSUMPTIONS to confirm before quoting")
t.add("CostWiring", "DC wiring, connectors and earthing", "₹ / piece", 120, FMT_INR,
      "ASSUMPTION: module leads, MC4 connectors, string cabling inside the lumber. About ₹3/Wp.", key=True)
t.add("CostInv", "Inverter, combiner, protection and metering", "₹ / Wp", 5.0, FMT_INR2,
      "ASSUMPTION. Wiring ₹3 + inverter ₹5 = ₹8/Wp, the earlier ROI model's whole electrical BOS.", key=True)
t.add("CostLabour", "Assembly and installation labour", "₹ / piece", 150, FMT_INR,
      "ASSUMPTION: bonding or fixing the glass, placing in the lumber, wiring, laying, commissioning. ₹600/m² against ₹300/m² in the earlier ROI model.", key=True)
t.add("MgmtPct", "Project management, QA and overhead at rollout", "% of direct cost", 0.08, FMT_PCT,
      "ASSUMPTION. Applied to the PV layer, electrical and labour at rollout. In the pilot these are the valued service lines in section 4 of 'Product & Costs'.", key=True)
t.add("MarginPct", "GridMats gross margin", "% of selling price", 0.25, FMT_PCT,
      "ASSUMPTION. The OC-LW 40 ladder price of ₹690/m² against ₹500 cost implies 27.5% on the base grid.", key=True)
t.add("VolDown", "Cost-down on module, glass and lumber at rollout volume", "% of today's cost", (0.0, 0.10, 0.20), FMT_PCT,
      "ASSUMPTION. Today's costs are pilot-lot costs. Conservative: none. Applies from the rollout threshold on 'Product & Costs'.", key=True)
t.add("GridMargin", "Margin the ladder price earns on the base grid (check)", "%", "=1-CostGrid*ModPerM2/GridT1", FMT_PCT, "1 − cost ÷ ladder price (under 1,000 m²).")
t.sec("3. Price factors and the PV-layer, electrical and labour prices that feed the model")
t.add("PFPilot", "Price factor, pilot: cost → price", "×", "=1/(1-MarginPct)", '0.000', "Price = cost ÷ (1 − margin).")
t.add("PFRoll", "Price factor, rollout: cost → price", "×", "=(1+MgmtPct)/(1-MarginPct)", '0.000', "Price = cost × (1 + management) ÷ (1 − margin).")
t.add("IsScale", "Rollout area is at or above the volume threshold (1 = yes)", "flag", "=IF(RO_Area>=ScaleFrom,1,0)", FMT_NUM, "Threshold on 'Product & Costs'.")
t.add("VDEff", "Volume cost-down applied at the rollout area", "%", "=VolDown*IsScale", FMT_PCT, "")
t.add("PFEff", "Price factor applied at the rollout area", "×", "=IF(IsScale=1,PFRoll,PFPilot)", '0.000', "")
t.add("PVMGPilot", "PV layer, pilot: modules + glass price", "₹ / m²", "=(CostModule+CostGlass)*ModPerM2*PFPilot", FMT_INR, "Feeds 'Pilot BoQ'.")
t.add("PVLWPilot", "PV layer, pilot: MLP lumber + wiring price", "₹ / m²", "=(CostLumber+CostWiring)*ModPerM2*PFPilot", FMT_INR, "Feeds 'Pilot BoQ'.")
t.add("PVPilot", "PV layer, pilot price", "₹ / m²", "=PVMGPilot+PVLWPilot", FMT_INR, "Module, glass, lumber and wiring.")
t.add("PVScale", "PV layer, rollout price", "₹ / m²", "=((CostModule+CostGlass+CostLumber)*(1-VolDown)+CostWiring)*ModPerM2*PFRoll", FMT_INR, "After the volume cost-down and management.")
t.add("BOSPilot", "Electrical, pilot price", "₹ / Wp", "=CostInv*PFPilot", FMT_INR2, "")
t.add("BOSScale", "Electrical, rollout price", "₹ / Wp", "=CostInv*PFRoll", FMT_INR2, "")
t.add("InstPilot", "Installation labour, pilot price", "₹ / m²", "=CostLabour*ModPerM2*PFPilot", FMT_INR, "")
t.add("InstScale", "Installation labour, rollout price", "₹ / m²", "=CostLabour*ModPerM2*PFRoll", FMT_INR, "")
t.add("PVPilotSqftCB", "PV layer, pilot price per sq ft", "₹ / sq ft", "=PVPilot/10.7639", FMT_INR, "₹ per sq ft first.")
t.add("PVScaleSqft", "PV layer, rollout price per sq ft", "₹ / sq ft", "=PVScale/10.7639", FMT_INR, "")

r = t.row + 1
section(ws, r, "4. Where the money goes (₹ per m² of walkway): GridMats cost against client price", ncols=7)
header(ws, r + 1, ["Component", "Unit", "Pilot: cost", "Pilot: price", "Rollout: cost", "Rollout: price", "Note"])
cb_rows = [
    ("Solar module", "CostModule*ModPerM2", "CostModule*ModPerM2*(1-VDEff)", "PF", "Priced with the PV layer."),
    ("12 mm toughened glass", "CostGlass*ModPerM2", "CostGlass*ModPerM2*(1-VDEff)", "PF", ""),
    ("MLP lumber", "CostLumber*ModPerM2", "CostLumber*ModPerM2*(1-VDEff)", "PF", ""),
    ("DC wiring and connectors", "CostWiring*ModPerM2", "CostWiring*ModPerM2", "PF", ""),
    ("GridMats OC-LW 40 base grid", "CostGrid*ModPerM2", "CostGrid*ModPerM2", "GRID", "Sold at the published ladder price."),
    ("Inverter, combiner, protection", "CostInv*PowerDens", "CostInv*PowerDens", "PF", ""),
    ("Installation labour", "CostLabour*ModPerM2", "CostLabour*ModPerM2", "PF", ""),
    ("Base works (subgrade, geotextile, WMM, M-sand bedding, edges)", "IF(BuildType=\"New walkway\",CivilNew,CivilOverlay)", "IF(BuildType=\"New walkway\",CivilNew,CivilOverlay)", "PASS", "Passed through at cost. WMM thickness follows the site CBR."),
    ("Logistics and freight", "LOG", "LOG", "LOGP", "% of hardware, on cost for the cost columns and on price for the price columns."),
    ("Grid integration and metering", "IntegLump/PilotArea", "RO_Integ", "PASS", "Passed through at cost."),
    ("Project management and overhead", "0", "MgmtPct*IsScale*(E{a}+E{b}+E{c}+E{d}+E{f}+E{g})", "MGMT", "Rollout only; in the pilot it is the valued service lines. Included in the rollout prices above."),
]
r0 = r + 2
for i, (lab, pc_, rc_, kind, nt) in enumerate(cb_rows):
    rr = r0 + i
    put(ws, f"A{rr}", lab, kind="label")
    put(ws, f"B{rr}", "₹ / m²", kind="note")
    note(ws, f"G{rr}", nt)
rowof = {i: r0 + i for i in range(len(cb_rows))}
for i, (lab, pc_, rc_, kind, nt) in enumerate(cb_rows):
    rr = rowof[i]
    hw_idx = [0, 1, 2, 3, 4, 5, 7]
    if pc_ == "LOG":
        pc_ = "LogPct*(" + "+".join(f"C{rowof[k]}" for k in hw_idx) + ")"
        rc_ = "LogPct*(" + "+".join(f"E{rowof[k]}" for k in hw_idx) + ")"
    if kind == "MGMT":
        rc_ = "MgmtPct*IsScale*(" + "+".join(f"E{rowof[k]}" for k in (0, 1, 2, 3, 5, 6)) + ")"
    put(ws, f"C{rr}", "=" + pc_, FMT_INR)
    put(ws, f"E{rr}", "=" + rc_, FMT_INR)
    if kind == "PF":
        put(ws, f"D{rr}", f"=C{rr}*PFPilot", FMT_INR)
        put(ws, f"F{rr}", f"=E{rr}*PFEff", FMT_INR)
    elif kind == "GRID":
        put(ws, f"D{rr}", "=GridT1", FMT_INR)
        put(ws, f"F{rr}", "=RO_Grid", FMT_INR)
    elif kind == "PASS":
        put(ws, f"D{rr}", f"=C{rr}", FMT_INR)
        put(ws, f"F{rr}", f"=E{rr}", FMT_INR)
    elif kind == "LOGP":
        put(ws, f"D{rr}", "=LogPct*(" + "+".join(f"D{rowof[k]}" for k in hw_idx) + ")", FMT_INR)
        put(ws, f"F{rr}", "=RO_Log", FMT_INR)
    elif kind == "MGMT":
        put(ws, f"D{rr}", 0, FMT_INR)
        put(ws, f"F{rr}", 0, FMT_INR)
rl = r0 + len(cb_rows) - 1
rt = rl + 1
put(ws, f"A{rt}", "Total per m²", kind="label", bold=True)
for c in "CDEF":
    put(ws, f"{c}{rt}", f"=SUM({c}{r0}:{c}{rl})", FMT_INR, bold=True)
    ws[f"{c}{rt}"].border = TOP2
rows_extra = [
    ("GridMats margin (₹ / m²)", f"=D{rt}-C{rt}", f"=F{rt}-E{rt}", FMT_INR),
    ("GridMats margin (% of price)", f"=(D{rt}-C{rt})/D{rt}", f"=(F{rt}-E{rt})/F{rt}", FMT_PCT),
    ("Price per sq ft", f"=D{rt}/10.7639", f"=F{rt}/10.7639", FMT_INR),
    ("Price per Wp (all-in, ex-services)", f"=D{rt}/PowerDens", f"=F{rt}/PowerDens", FMT_INR2),
    ("Cost per Wp (GridMats)", f"=C{rt}/PowerDens", f"=E{rt}/PowerDens", FMT_INR2),
]
for i, (lab, fp, fr, fmt) in enumerate(rows_extra):
    rr = rt + 1 + i
    put(ws, f"A{rr}", lab, kind="label")
    put(ws, f"D{rr}", fp, fmt)
    put(ws, f"F{rr}", fr, fmt)
dname("CB_CostPilotM2", "Cost Build-up", f"C{rt}")
dname("CB_PricePilotM2", "Cost Build-up", f"D{rt}")
dname("CB_CostRollM2", "Cost Build-up", f"E{rt}")
dname("CB_PriceRollM2", "Cost Build-up", f"F{rt}")
dname("CB_MarginPilotPct", "Cost Build-up", f"D{rt+2}")
dname("CB_MarginRollPct", "Cost Build-up", f"F{rt+2}")
chk = rt + 1 + len(rows_extra)
put(ws, f"A{chk}", "Check: price columns agree with 'Pilot BoQ' and 'Rollout Economics'", kind="label")
put(ws, f"D{chk}", f"=IF(AND(ABS(D{rt}-PilotPhysM2)<1,ABS(F{rt}-RO_CostM2)<1),\"OK\",\"CHECK\")", None, bold=True, align="center")

r = chk + 2
section(ws, r, "5. GridMats' own money in the pilot", ncols=7)
gm = [
    ("Hardware and installation at GridMats cost (₹)", "=CB_CostPilotM2*PilotArea", FMT_INR, "GM_PilotCost"),
    ("Valued services given free: design, supervision, travel, IoT (₹)", "=PilotList-PilotPhys", FMT_INR, "GM_PilotServices"),
    ("Total pilot cost to GridMats (₹)", "=GM_PilotCost+GM_PilotServices", FMT_INR, "GM_PilotTotal"),
    ("Client pays, ex-GST (₹)", "=PilotPays", FMT_INR, None),
    ("GridMats' net investment in the pilot (₹)", "=GM_PilotTotal-PilotPays", FMT_INR, "GM_PilotNet"),
]
for i, (lab, f, fmt, nm) in enumerate(gm):
    rr = r + 1 + i
    put(ws, f"A{rr}", lab, kind="label", bold=nm == "GM_PilotNet")
    put(ws, f"D{rr}", f, fmt, bold=nm == "GM_PilotNet")
    if nm:
        dname(nm, "Cost Build-up", f"D{rr}")
r = r + len(gm) + 2

section(ws, r, "6. What it takes for the campus rollout to pay back", ncols=7)
t.row = r + 1
t.add("TargetPB", "Target payback for this view", "years", 6, FMT_NUM1, "ASSUMPTION: a common corporate hurdle for green capex. Change to the client's.", key=True)
t.add("BE_Save", "Year-1 electricity saving per m² of walkway", "₹ / m²", "=PowerDens/1000*SpecYield*Tariff", FMT_INR, "Rated power × specific yield × tariff.")
t.add("BE_Price", "Highest installed price that meets the target, net of the paver credit", "₹ / m²", "=(TargetPB*BE_Save+RO_Credit)/(1+TargetPB*OMPct)", FMT_INR, "Solves (price − credit) ÷ (saving − O&M × price) = target payback.")
t.add("BE_Now", "Rollout installed price now", "₹ / m²", "=RO_CostM2", FMT_INR, "From 'Rollout Economics'.")
t.add("BE_Gap", "Price reduction needed (negative = target already met)", "%", "=1-BE_Price/BE_Now", FMT_PCT, "")
t.add("BE_PV", "PV-layer price that meets the target, other items unchanged", "₹ / m²", "=(BE_Price-RO_Integ-RO_Inst)/(1+LogPct)-RO_Civil-RO_Grid-RO_BOS", FMT_INR, "Backs the PV layer out of the installed price.")
t.add("BE_Today", "Module + glass + lumber cost per piece today", "₹ / piece", "=CostModule+CostGlass+CostLumber", FMT_INR, "Before any volume cost-down.")
t.add("BE_Piece", "Module + glass + lumber cost per piece that meets the target", "₹ / piece", "=MAX(0,((BE_PV/PFEff)/ModPerM2-CostWiring)/(1-VDEff))", FMT_INR,
      "Holding wiring, electrical, labour, margin and management as assumed.")
t.add("BE_Cut", "Cost reduction needed on module + glass + lumber", "%", "=1-BE_Piece/BE_Today", FMT_PCT, "")
r = t.row + 1
put(ws, f"A{r}", "Cost-down path: rollout price and payback as module, glass and lumber cost falls", kind="label", bold=True)
header(ws, r + 1, ["Cost vs today (before volume cost-down)", "Per piece (₹)", "Rollout installed (₹/sq ft)", "Simple payback (yrs)", "Net of paver credit (₹/sq ft)"])
for i, k in enumerate([1.0, 0.9, 0.8, 0.7, 0.6, 0.5]):
    rr = r + 2 + i
    put(ws, f"A{rr}", k, FMT_PCT, kind="input", bold=True, align="center")
    put(ws, f"B{rr}", f"=BE_Today*A{rr}", FMT_INR)
    newc = f"(RO_CostM2-(1-A{rr})*BE_Today*ModPerM2*(1-VDEff)*PFEff*(1+LogPct))"
    put(ws, f"C{rr}", f"={newc}/10.7639", FMT_INR)
    put(ws, f"D{rr}", f"=IFERROR(({newc}-RO_Credit)/(BE_Save-OMPct*{newc}),\"n/a\")", '0.0', align="center")
    put(ws, f"E{rr}", f"=({newc}-RO_Credit)/10.7639", FMT_INR)
rng = f"D{r+2}:D{r+7}"
from openpyxl.formatting.rule import CellIsRule as _CIR
ws.conditional_formatting.add(rng, _CIR(operator="lessThanOrEqual", formula=["6"], fill=PatternFill("solid", fgColor="E3F1E6")))
ws.conditional_formatting.add(rng, _CIR(operator="between", formula=["6.0001", "9"], fill=PatternFill("solid", fgColor="FFF4D6")))
ws.conditional_formatting.add(rng, _CIR(operator="greaterThan", formula=["9"], fill=PatternFill("solid", fgColor="FBEAEC")))
ws.freeze_panes = "A5"

# ================================================================ IRRADIANCE
ws = irr
title(ws, "Sunshine library: monthly global horizontal irradiance (kWh/m²/day)",
      "Pick the city on 'Client'. Values are typical-year climatology; replace with site data (Global Solar Atlas / NASA POWER) after the site walk.")
widths(ws, {"A": 4, "B": 16, "O": 14, "P": 12, "Q": 70})
for i in range(3, 15):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 7.5
header(ws, 5, ["", "City"] + MONTHS + ["kWh/m²/yr", "Avg / day", "Source and confidence"])
put(ws, "B6", "Days in month", kind="label", bold=True)
for j, dd in enumerate(DAYS):
    put(ws, f"{openpyxl.utils.get_column_letter(3+j)}6", dd, FMT_NUM, kind="calc", align="center")
CITIES = [
    ("Mumbai", [4.8, 5.6, 6.3, 6.7, 6.7, 4.6, 3.7, 3.8, 4.5, 5.1, 4.8, 4.6], "Calibrated to NASA-derived Navi Mumbai data (annual mean 5.16, Jan 4.74, Jul 3.68 kWh/m²/day) and Meteonorm 8.2 (1,814 kWh/m²/yr). Earlier ROI model: 5.3 average."),
    ("Bengaluru", [5.4, 6.2, 6.6, 6.5, 6.0, 4.8, 4.3, 4.6, 5.0, 4.8, 4.6, 4.9], "Indicative typical year."),
    ("Chennai", [5.0, 6.0, 6.6, 6.7, 6.4, 5.6, 5.1, 5.3, 5.4, 4.6, 4.0, 4.1], "Indicative typical year."),
    ("Hyderabad", [5.0, 5.8, 6.4, 6.7, 6.6, 5.2, 4.5, 4.6, 4.9, 4.9, 4.7, 4.6], "Indicative typical year."),
    ("Delhi NCR", [3.6, 4.6, 5.8, 6.7, 7.0, 6.4, 5.4, 5.1, 5.3, 4.9, 4.0, 3.4], "Indicative typical year."),
    ("Pune", [5.2, 6.0, 6.6, 6.8, 6.7, 4.8, 3.8, 4.0, 4.7, 5.0, 4.9, 4.8], "Indicative typical year."),
    ("Ahmedabad", [4.6, 5.4, 6.3, 6.9, 7.1, 6.2, 4.8, 4.6, 5.4, 5.3, 4.6, 4.3], "Indicative typical year."),
    ("Jamnagar", [4.8, 5.6, 6.5, 7.0, 7.2, 6.3, 4.9, 4.8, 5.6, 5.5, 4.8, 4.5], "Indicative typical year."),
    ("Kolkata", [4.2, 4.9, 5.6, 6.1, 6.0, 4.8, 4.3, 4.4, 4.3, 4.4, 4.3, 4.0], "Indicative typical year."),
]
for i, (city, vals, srcnote) in enumerate(CITIES):
    r = 7 + i
    put(ws, f"B{r}", city, kind="label", bold=True)
    for j, v in enumerate(vals):
        put(ws, f"{openpyxl.utils.get_column_letter(3+j)}{r}", v, FMT_NUM1, kind="input")
    put(ws, f"O{r}", f"=SUMPRODUCT(C{r}:N{r},$C$6:$N$6)", FMT_NUM)
    put(ws, f"P{r}", f"=O{r}/365", FMT_NUM2)
    note(ws, f"Q{r}", srcnote)
IRR_ROW0 = 7


# ================================================================ MODULE TECH
ws = mt
title(ws, "Module technology library (2026)",
      "Pick the technology on 'Product & Costs'. TOPCon is now the mainstream cell; PERC is being phased out; Reliance makes HJT modules at Jamnagar.")
widths(ws, {"A": 3, "B": 22, "C": 13, "D": 30, "E": 13, "F": 13, "G": 12, "H": 12, "I": 14, "J": 44, "K": 60})
header(ws, 5, ["", "Technology", "Cell efficiency (mass prod.)", "Best cells (2026)", "Module efficiency", "Temp. coefficient (%/°C)",
               "Year-1 degradation", "Annual degradation", "Module price, India (₹/Wp)", "Notes", "Sources"])
TECH = [
    ("Mono PERC (p-type)", 0.233, "Mature; no longer advancing", 0.212, -0.0035, 0.02, 0.005, 13.8,
     "Being phased out; TOPCon premium over PERC has collapsed to about 0-10%.", "JMK Research Apr 2026 (ex-works, 500 Wp); 2026 technology comparisons."),
    ("TOPCon (n-type)", 0.258, "26.66% industrial M10 cell (JinkoSolar, Feb 2026); 27.79% lab cell (ISFH-certified)", 0.235, -0.0029, 0.01, 0.004, 14.2,
     "About two-thirds of global capacity; commercial modules 22.5-24.8%, best 25.91%. DCR TOPCon ₹22.3/Wp.",
     "pv-magazine (25 Feb 2026); taiyangnews; Canadian Solar TOPHiKu6 and Jinko Tiger Neo datasheets (−0.29%/°C, 1% then 0.4%/yr); JMK Apr 2026."),
    ("HJT (n-type)", 0.258, "Lowest temperature loss of the three", 0.240, -0.0025, 0.01, 0.003, None,
     "Commercial modules 23-25%; price premium over TOPCon. Reliance produces HJT modules at Jamnagar.",
     "2026 technology comparisons (−0.24 to −0.26%/°C; 0.25-0.375%/yr)."),
]
for i, (nm, ce, best, me, tc, d1, da, price, notes, srcs) in enumerate(TECH):
    r = 6 + i
    put(ws, f"B{r}", nm, kind="label", bold=True)
    put(ws, f"C{r}", ce, FMT_PCT, kind="input")
    note(ws, f"D{r}", best)
    put(ws, f"E{r}", me, FMT_PCT, kind="input")
    put(ws, f"F{r}", tc, '0.00%', kind="input")
    put(ws, f"G{r}", d1, FMT_PCT, kind="input")
    put(ws, f"H{r}", da, FMT_PCT, kind="input")
    put(ws, f"I{r}", price if price is not None else "premium", FMT_INR2 if price else None, kind="input")
    note(ws, f"J{r}", notes)
    note(ws, f"K{r}", srcs)
    ws.row_dimensions[r].height = 45
section(ws, 11, "What each technology means for this walkway (selected scenario; pilot size, same rated power)", ncols=11)
header(ws, 12, ["", "Technology", "Temperature loss", "Performance ratio", "Specific yield (kWh/kWp/yr)", "Pilot year-1 kWh",
                "Pilot 20-year kWh", "Technology-derived module (W)", "Walkway W/m² (derived)", "Reading", ""])
for i in range(3):
    r = 13 + i
    src_r = 6 + i
    put(ws, f"B{r}", f"=B{src_r}", kind="link", bold=True)
    put(ws, f"C{r}", f"=-F{src_r}*(Tcell-25)", FMT_PCT)
    put(ws, f"D{r}", f"=(1-C{r})*(1-LossSoil)*(1-LossFoot)*(1-LossRefl)*(1-LossMism)*(1-LossInv)*(1-LossAvail)", FMT_PCT)
    put(ws, f"E{r}", f"=SunYear*D{r}*(1-ShadeLoss)", FMT_NUM)
    put(ws, f"F{r}", f"=E{r}*PilotKWp*(1-G{src_r})", FMT_NUM)
    put(ws, f"G{r}", f"=F{r}*SUMPRODUCT((1-H{src_r}-WearAdder)^(ROW($A$1:$A$20)-1))", FMT_NUM)
    put(ws, f"H{r}", f"=ModArea*E{src_r}*(1-WalkPackLoss)*1000", FMT_NUM1)
    put(ws, f"I{r}", f"=H{r}*ModPerM2", FMT_NUM)
    dname(f"TC_PR_{i}", "Module Tech", f"D{r}")
    dname(f"TC_Y1_{i}", "Module Tech", f"F{r}")
    dname(f"TC_Y20_{i}", "Module Tech", f"G{r}")
    dname(f"TC_WM2_{i}", "Module Tech", f"I{r}")
    dname(f"TC_TL_{i}", "Module Tech", f"C{r}")
note(ws, "J13", "PERC loses the most to heat on a hot, flat walkway.")
note(ws, "J14", "TOPCon: the default. Lower heat loss and slower ageing than PERC at a similar price.")
note(ws, "J15", "HJT: lowest heat loss; worth testing with Reliance's own Jamnagar cells.")
note(ws, "B17", "Pilot year-1 kWh includes first-year degradation. The derived module rating is an upgrade path to be validated; the client price is quoted on the 40 W design basis.")
ws.merge_cells("B17:K18")

# ================================================================ ENERGY
ws = en
title(ws, "Energy: what the walkway generates", "Monthly generation for the pilot at the chosen city, after performance ratio and site shading.")
widths(ws, {"A": 12, "B": 10, "C": 14, "D": 16, "E": 18, "F": 16, "G": 16})
section(ws, 4, "Pilot size", ncols=7)
for i, (lab, f, fmt, nm) in enumerate([
        ("Pilot area (m²)", "=PilotArea", FMT_NUM1, None),
        ("Modules", "=PilotArea*ModPerM2", FMT_NUM, "PilotModules"),
        ("Rated power (kWp)", "=PilotArea*PowerDens/1000", FMT_NUM2, "PilotKWp"),
        ("Sunshine profile", "=IrrCity", None, None)]):
    put(ws, f"A{5+i}", lab, kind="label")
    put(ws, f"D{5+i}", f, fmt, kind="link" if nm is None else "calc")
    if nm:
        dname(nm, "Energy", f"D{5+i}")
header(ws, 10, ["Month", "Days", "Sunshine (kWh/m²/day)", "Sunshine (kWh/m²/month)", "Pilot generation (kWh)", "Per m² (kWh)", "Per day (kWh)"])
for i, m in enumerate(MONTHS):
    r = 11 + i
    colL = openpyxl.utils.get_column_letter(3 + i)
    put(ws, f"A{r}", m, kind="label")
    put(ws, f"B{r}", DAYS[i], FMT_NUM)
    put(ws, f"C{r}", f"=INDEX(Irradiance!{colL}${IRR_ROW0}:{colL}${IRR_ROW0+len(CITIES)-1},MATCH(IrrCity,Irradiance!$B${IRR_ROW0}:$B${IRR_ROW0+len(CITIES)-1},0))", FMT_NUM2, kind="link")
    put(ws, f"D{r}", f"=C{r}*B{r}", FMT_NUM1)
    put(ws, f"E{r}", f"=D{r}*PilotKWp*PR*(1-ShadeLoss)", FMT_NUM)
    put(ws, f"F{r}", f"=E{r}/PilotArea", FMT_NUM1)
    put(ws, f"G{r}", f"=E{r}/B{r}", FMT_NUM2)
r = 23
put(ws, f"A{r}", "Year 1*", kind="label", bold=True)
for c in "BDEF":
    put(ws, f"{c}{r}", f"=SUM({c}11:{c}22)", FMT_NUM, bold=True)
    ws[f"{c}{r}"].border = TOP2
put(ws, f"G{r}", f"=E{r}/365", FMT_NUM2, bold=True)
dname("SunYear", "Energy", f"D{r}")
dname("PilotKWh1", "Energy", f"E{r}")
dname("KWhPerM2", "Energy", f"F{r}")
dname("PilotKWhDay", "Energy", f"G{r}")
note(ws, f"A{r+1}", "* Before first-year degradation; the cash-flow sheets apply it.")
ws.merge_cells(f"A{r+1}:G{r+1}")
r = 25
section(ws, r, "Key energy figures", ncols=7)
kef = [("Specific yield (kWh per kWp per year)", "=SunYear*PR*(1-ShadeLoss)", FMT_NUM, "SpecYield"),
       ("Generation per m² per day, average (kWh)", "=KWhPerM2/365", FMT_NUM2, "KWhPerM2Day"),
       ("Best month (kWh)", "=MAX(E11:E22)", FMT_NUM, "BestMonth"),
       ("Weakest month (kWh)", "=MIN(E11:E22)", FMT_NUM, "WorstMonth"),
       ("CO2 avoided in year 1 (kg)", "=PilotKWh1*GridEF", FMT_NUM, "PilotCO2kg"),
       ("Recycled PP in the pilot base (kg)", "=PilotArea*PlasticLW", FMT_NUM, "PilotPPkg")]
for i, (lab, f, fmt, nm) in enumerate(kef):
    put(ws, f"A{r+1+i}", lab, kind="label")
    ws.merge_cells(f"A{r+1+i}:C{r+1+i}")
    put(ws, f"D{r+1+i}", f, fmt, bold=True)
    dname(nm, "Energy", f"D{r+1+i}")

# ================================================================ PILOT BOQ
ws = boq
title(ws, "Pilot bill of quantities and offer",
      "List = GridMats standard value. 'GridMats contributes' = what the pilot offer covers. Ex-GST unless stated.")
widths(ws, {"A": 62, "B": 11, "C": 9, "D": 13, "E": 14, "F": 15, "G": 15, "H": 48})
header(ws, 4, ["Item", "Qty", "Unit", "Rate (₹)", "List value (₹)", "GridMats contributes (₹)", "Client pays (₹)", "Note"])
items = [
    ("Site assessment, shading study and engineering design", "1", "lot", "=DesignLump", "=E{r}", "Included", "svc"),
    ("=IF(BuildType=\"New walkway\",\"Base works: compacted subgrade, geotextile, \"&WMMThick&\" mm WMM, \"&MsandThick&\" mm M-sand bedding, edges\",\"Base works: overlay levelling and edge ramps\")", "=PilotArea", "m²",
     "=IF(BuildType=\"New walkway\",CivilNew,CivilOverlay)", "0", "", "hw"),
    ("GridMats OC-LW 40 structural base (100% recycled PP), 4 grids per m²", "=PilotArea", "m²", "=GridT1", "0", "", "hw"),
    ("Walkable PV layer: 40 W solar modules under 12 mm toughened glass, 4 per m²", "=PilotArea", "m²", "=PVMGPilot", "=E{r}*CoFund", "Co-funded by GridMats", "hw"),
    ("MLP lumber frame (40 × 40 mm) and module wiring", "=PilotArea", "m²", "=PVLWPilot", "=E{r}*CoFund", "Co-funded by GridMats", "hw"),
    ("Electrical: inverter, DC/AC cabling, protection, earthing", "=PilotKWp*1000", "Wp", "=BOSPilot", "0", "", "hw"),
    ("Logistics and freight", "=SUM(E6:E10)", "₹", "=LogPct", "0", "% of hardware", "log"),
    ("Grid integration and metering", "1", "lot", "=IntegLump", "0", "", "fix"),
    ("Installation labour", "=PilotArea", "m²", "=InstPilot", "0", "Built by teams GridMats trains", "fix"),
    ("GridMats supervision and training on site", "=SupDays", "days", "=ManDay", "=E{r}", "Included", "svc"),
    ("Supervisor travel and stay", "1", "lot", "=TravelLump", "=E{r}", "Included", "svc"),
    ("IoT sensors, live dashboard and performance study", "1", "lot", "=IoTPkg", "=E{r}", "Included", "svc"),
]
r = 5
first = r
for lab, qty, unit, rate, contrib, nt, kind in items:
    put(ws, f"A{r}", lab, kind="label")
    put(ws, f"B{r}", qty if qty.startswith("=") else int(qty), FMT_NUM1 if unit != "₹" else FMT_INR,
        kind="calc" if qty.startswith("=") else "input")
    put(ws, f"C{r}", unit, kind="note")
    put(ws, f"D{r}", rate, FMT_PCT if kind == "log" else (FMT_INR2 if unit == "Wp" else FMT_INR), kind="link")
    put(ws, f"E{r}", f"=B{r}*D{r}", FMT_INR)
    put(ws, f"F{r}", "=" + contrib.format(r=r).lstrip("=") if contrib != "0" else 0, FMT_INR)
    put(ws, f"G{r}", f"=E{r}-F{r}", FMT_INR)
    put(ws, f"H{r}", nt, kind="note")
    r += 1
last = r - 1
put(ws, f"A{r}", "Total (ex-GST)", kind="label", bold=True)
for c in "EFG":
    put(ws, f"{c}{r}", f"=SUM({c}{first}:{c}{last})", FMT_INR, bold=True)
    ws[f"{c}{r}"].border = TOP2
dname("PilotList", "Pilot BoQ", f"E{r}")
dname("PilotContrib", "Pilot BoQ", f"F{r}")
dname("PilotPays", "Pilot BoQ", f"G{r}")
r += 1
put(ws, f"A{r}", "GST on what the client pays", kind="label")
put(ws, f"D{r}", "=GSTRate", FMT_PCT, kind="link")
put(ws, f"G{r}", f"=G{r-1}*D{r}", FMT_INR)
r += 1
put(ws, f"A{r}", "Client pays, including GST", kind="label", bold=True)
put(ws, f"G{r}", f"=G{r-2}+G{r-1}", FMT_INR, bold=True)
dname("PilotPaysGST", "Pilot BoQ", f"G{r}")
r += 2
section(ws, r, "Summary", ncols=8)
summ = [
    ("Physical installed cost of the pilot, list (hardware, logistics, integration, installation)", f"=SUMPRODUCT(E{first}:E{last},--(H{first}:H{last}<>\"Included\"))", FMT_INR, "PilotPhys"),
    ("Physical installed cost per sq ft (₹)", "=PilotPhys/(PilotArea*10.7639)", FMT_INR, "PilotPhysSqft"),
    ("Physical installed cost per m² (₹)", "=PilotPhys/PilotArea", FMT_INR, "PilotPhysM2"),
    ("Physical installed cost per kWp (₹)", "=PilotPhys/PilotKWp", FMT_INR, "PilotPhysKWp"),
    ("Client pays per sq ft, ex-GST (₹)", "=PilotPays/(PilotArea*10.7639)", FMT_INR, "PilotPaysSqft"),
    ("GridMats contribution as a share of list", "=PilotContrib/PilotList", FMT_PCT, "PilotContribPct"),
    ("PV layer price per sq ft, pilot (₹)", "=PVPilot/10.7639", FMT_INR, "PVPilotSqft"),
]
for i, (lab, f, fmt, nm) in enumerate(summ):
    put(ws, f"A{r+1+i}", lab, kind="label")
    put(ws, f"G{r+1+i}", f, fmt, bold=True)
    dname(nm, "Pilot BoQ", f"G{r+1+i}")
ws.freeze_panes = "A5"


# ================================================================ cash-flow helper (years down the rows)
def cashflow(ws, sheet, r0, capex_name, gen1, om_base, label, inv_kwp, extra_cols=None):
    """Write a 20-year table at row r0. Returns dict of key cell refs."""
    hdr = ["Year", "Degradation factor", "Generation (kWh)", "Tariff (₹/kWh)", "Savings (₹)", "O&M (₹)", "Inverter (₹)",
           "Net cash flow (₹)", "Cumulative (₹)", "Discount factor", "Discounted CF (₹)", "Cumulative discounted (₹)"]
    header(ws, r0, hdr)
    put(ws, f"A{r0+1}", 0, FMT_NUM, align="center")
    put(ws, f"H{r0+1}", f"=-{capex_name}", FMT_INR)
    put(ws, f"I{r0+1}", f"=H{r0+1}", FMT_INR)
    put(ws, f"J{r0+1}", 1, '0.000')
    put(ws, f"K{r0+1}", f"=H{r0+1}*J{r0+1}", FMT_INR)
    put(ws, f"L{r0+1}", f"=K{r0+1}", FMT_INR)
    for y in range(1, 21):
        r = r0 + 1 + y
        put(ws, f"A{r}", y, FMT_NUM, align="center")
        put(ws, f"B{r}", f"=IF(A{r}<=Life,(1-Deg1)*(1-Degr)^(A{r}-1),0)", '0.000')
        put(ws, f"C{r}", f"={gen1}*B{r}", FMT_NUM)
        put(ws, f"D{r}", f"=Tariff*(1+TariffEsc)^(A{r}-1)", FMT_INR2)
        put(ws, f"E{r}", f"=C{r}*D{r}", FMT_INR)
        put(ws, f"F{r}", f"=IF(A{r}<=Life,-OMPct*{om_base},0)", FMT_INR)
        put(ws, f"G{r}", f"=IF(AND(A{r}=InvYear,InvYear<Life),-InvRepl*{inv_kwp},0)", FMT_INR)
        put(ws, f"H{r}", f"=E{r}+F{r}+G{r}", FMT_INR)
        put(ws, f"I{r}", f"=I{r-1}+H{r}", FMT_INR)
        put(ws, f"J{r}", f"=1/(1+DiscRate)^A{r}", '0.000')
        put(ws, f"K{r}", f"=H{r}*J{r}", FMT_INR)
        put(ws, f"L{r}", f"=L{r-1}+K{r}", FMT_INR)
    a, b = r0 + 1, r0 + 21
    m = {}
    rr = b + 2
    section(ws, rr, f"Results: {label}", ncols=12)
    res = [
        ("Year-1 generation (kWh)", f"=C{a+1}", FMT_NUM, "Gen1"),
        ("Year-1 savings (₹)", f"=E{a+1}", FMT_INR, "Save1"),
        ("Year-1 O&M (₹)", f"=-F{a+1}", FMT_INR, "OM1"),
        ("Simple payback (years)",
         f"=IF(COUNTIF(I{a+1}:I{b},\"<0\")>=20,\"> 20\",COUNTIF(I{a+1}:I{b},\"<0\")+(-INDEX(I{a}:I{b},COUNTIF(I{a+1}:I{b},\"<0\")+1))/INDEX(H{a}:H{b},COUNTIF(I{a+1}:I{b},\"<0\")+2))",
         FMT_YRS, "Payback"),
        ("Discounted payback (years)",
         f"=IF(COUNTIF(L{a+1}:L{b},\"<0\")>=20,\"> 20\",COUNTIF(L{a+1}:L{b},\"<0\")+(-INDEX(L{a}:L{b},COUNTIF(L{a+1}:L{b},\"<0\")+1))/INDEX(K{a}:K{b},COUNTIF(L{a+1}:L{b},\"<0\")+2))",
         FMT_YRS, "DPayback"),
        ("NPV over 20 years (₹)", f"=L{b}", FMT_INR, "NPV"),
        ("IRR over 20 years", f"=IFERROR(IRR(H{a}:H{b}),\"n/a\")", FMT_PCT, "IRR"),
        ("20-year net cash (₹)", f"=I{b}", FMT_INR, "Cum20"),
        ("Return on investment over 20 years (×)", f"=IFERROR(I{b}/{capex_name},\"n/a\")", FMT_X, "ROI"),
    ]
    for i, (lab, f, fmt, key) in enumerate(res):
        put(ws, f"A{rr+1+i}", lab, kind="label")
        ws.merge_cells(f"A{rr+1+i}:D{rr+1+i}")
        put(ws, f"E{rr+1+i}", f, fmt, bold=True)
        m[key] = f"E{rr+1+i}"
    m["rows"] = (a, b)
    m["end"] = rr + 1 + len(res)
    return m


# ================================================================ PILOT ECONOMICS
ws = pe
title(ws, "Pilot economics", "Two views: on what the client actually pays (with GridMats co-funding), and on the full physical cost (no co-funding).")
widths(ws, {"A": 8, "B": 12, "C": 14, "D": 12, "E": 14, "F": 12, "G": 12, "H": 15, "I": 15, "J": 11, "K": 15, "L": 16})
put(ws, "A4", "Client's pilot outlay, ex-GST (₹)", kind="label", bold=True)
put(ws, "H4", "=PilotPays", FMT_INR, kind="link", bold=True)
put(ws, "A5", "Full physical installed cost, ex-GST (₹)", kind="label")
put(ws, "H5", "=PilotPhys", FMT_INR, kind="link")
put(ws, "A6", "Year-1 generation (kWh)", kind="label")
put(ws, "H6", "=PilotKWh1", FMT_NUM, kind="link")
note(ws, "A7", "O&M is charged on the physical installed cost in both views. GST is excluded on the assumption that the client claims input tax credit.")
ws.merge_cells("A7:L7")
section(ws, 9, "A. On the client's outlay (pilot offer)", ncols=12)
m1 = cashflow(ws, "Pilot Economics", 10, "PilotPays", "PilotKWh1", "PilotPhys", "client's outlay", "PilotKWp")
for k, nm in [("Payback", "PE_Payback"), ("NPV", "PE_NPV"), ("IRR", "PE_IRR"), ("Save1", "PE_Save1"), ("OM1", "PE_OM1"), ("Cum20", "PE_Cum20"), ("DPayback", "PE_DPayback")]:
    dname(nm, "Pilot Economics", m1[k])
r0 = m1["end"] + 2
section(ws, r0, "B. On the full physical installed cost (no co-funding)", ncols=12)
m2 = cashflow(ws, "Pilot Economics", r0 + 1, "PilotPhys", "PilotKWh1", "PilotPhys", "full physical cost", "PilotKWp")
for k, nm in [("Payback", "PF_Payback"), ("NPV", "PF_NPV"), ("IRR", "PF_IRR"), ("Cum20", "PF_Cum20")]:
    dname(nm, "Pilot Economics", m2[k])

# ================================================================ ROLLOUT ECONOMICS
ws = ro
title(ws, "Rollout economics: campus walkway at rollout prices", "Uses the campus rollout area from 'Client'. No co-funding at rollout.")
widths(ws, {"A": 8, "B": 12, "C": 14, "D": 12, "E": 14, "F": 12, "G": 12, "H": 15, "I": 15, "J": 11, "K": 15, "L": 16, "M": 40})
section(ws, 4, "Installed cost per m² at the rollout area", ncols=12)
ro_rows = [
    ("Rollout area (m²)", "=CampusArea", FMT_NUM, "RO_Area"),
    ("Rated power (kWp)", "=RO_Area*PowerDens/1000", FMT_NUM1, "RO_KWp"),
    ("Base works (₹/m²)", "=IF(BuildType=\"New walkway\",CivilNew,CivilOverlay)", FMT_INR, "RO_Civil"),
    ("GridMats OC-LW 40 at this area (₹/m²)", "=IF(RO_Area<1000,GridT1,IF(RO_Area<2500,GridT2,IF(RO_Area<5000,GridT3,GridT4)))", FMT_INR, "RO_Grid"),
    ("Walkable PV layer (₹/m²)", "=IF(RO_Area>=ScaleFrom,PVScale,PVPilot)", FMT_INR, "RO_PV"),
    ("Electrical BOS (₹/m²)", "=IF(RO_Area>=ScaleFrom,BOSScale,BOSPilot)*PowerDens", FMT_INR, "RO_BOS"),
    ("Logistics (₹/m²)", "=LogPct*(RO_Civil+RO_Grid+RO_PV+RO_BOS)", FMT_INR, "RO_Log"),
    ("Grid integration (₹/m²)", "=ROUNDUP(RO_KWp/ConnKWp,0)*IntegLump/RO_Area", FMT_INR, "RO_Integ"),
    ("Installation (₹/m²)", "=IF(RO_Area>=ScaleFrom,InstScale,InstPilot)", FMT_INR, "RO_Inst"),
    ("Installed cost (₹/m²)", "=SUM(H7:H13)", FMT_INR, "RO_CostM2"),
    ("Installed cost (₹/sq ft)", "=RO_CostM2/10.7639", FMT_INR, "RO_CostSqft"),
    ("Less: walkway it replaces (₹/m², new build only)", "=IF(BuildType=\"New walkway\",ConvWalk,0)", FMT_INR, "RO_Credit"),
    ("Net extra cost for a generating walkway (₹/m²)", "=RO_CostM2-RO_Credit", FMT_INR, "RO_NetM2"),
    ("Total installed cost (₹)", "=RO_CostM2*RO_Area", FMT_INR, "RO_Capex"),
    ("Net investment after walkway credit (₹)", "=RO_NetM2*RO_Area", FMT_INR, "RO_NetCapex"),
    ("Year-1 generation (kWh)", "=RO_KWp*SpecYield", FMT_NUM, "RO_Gen1"),
    ("Installed cost per kWp, solar part (₹)", "=(RO_PV+RO_BOS+RO_Log+RO_Integ)/(PowerDens/1000)", FMT_INR, "RO_PerKWp"),
]
for i, (lab, f, fmt, nm) in enumerate(ro_rows):
    r = 5 + i
    put(ws, f"A{r}", lab, kind="label", bold=nm in ("RO_CostM2", "RO_NetM2"))
    ws.merge_cells(f"A{r}:G{r}")
    put(ws, f"H{r}", f, fmt, bold=nm in ("RO_CostM2", "RO_NetM2"))
    dname(nm, "Rollout Economics", f"H{r}")
m3 = cashflow(ws, "Rollout Economics", 24, "RO_NetCapex", "RO_Gen1", "RO_Capex", "campus rollout", "RO_KWp")
for k, nm in [("Payback", "RO_Payback"), ("NPV", "RO_NPV"), ("IRR", "RO_IRR"), ("Save1", "RO_Save1"), ("OM1", "RO_OM1"), ("Cum20", "RO_Cum20"), ("DPayback", "RO_DPayback")]:
    dname(nm, "Rollout Economics", m3[k])
a, b = m3["rows"]
r = m3["end"] + 2
section(ws, r, "After-tax view with 40% accelerated depreciation on the solar part", ncols=12)
header(ws, r + 1, ["Year", "Opening WDV (₹)", "Depreciation (₹)", "Tax shield (₹)", "After-tax cash flow (₹)"])
solar_part = "(RO_PV+RO_BOS+RO_Log+RO_Integ)*RO_Area"
for y in range(0, 21):
    rr = r + 2 + y
    put(ws, f"A{rr}", y, FMT_NUM, align="center")
    if y == 0:
        put(ws, f"B{rr}", 0, FMT_INR); put(ws, f"C{rr}", 0, FMT_INR); put(ws, f"D{rr}", 0, FMT_INR)
        put(ws, f"E{rr}", f"=H{a}", FMT_INR)
    else:
        put(ws, f"B{rr}", f"={solar_part}" if y == 1 else f"=B{rr-1}-C{rr-1}", FMT_INR)
        put(ws, f"C{rr}", f"=B{rr}*ADRate", FMT_INR)
        put(ws, f"D{rr}", f"=C{rr}*TaxRate", FMT_INR)
        put(ws, f"E{rr}", f"=H{a+y}*(1-TaxRate)+D{rr}", FMT_INR)
rr = r + 23
put(ws, f"A{rr}", "After-tax IRR", kind="label", bold=True)
ws.merge_cells(f"A{rr}:D{rr}")
put(ws, f"E{rr}", f"=IFERROR(IRR(E{r+2}:E{r+22}),\"n/a\")", FMT_PCT, bold=True)
dname("RO_IRRAT", "Rollout Economics", f"E{rr}")
rr += 1
put(ws, f"A{rr}", "Levelised cost of the walkway's power (₹/kWh)", kind="label", bold=True)
ws.merge_cells(f"A{rr}:D{rr}")
put(ws, f"E{rr}", f"=({solar_part}+SUMPRODUCT(-(F{a+1}:F{b}+G{a+1}:G{b}),J{a+1}:J{b}))/SUMPRODUCT(C{a+1}:C{b},J{a+1}:J{b})", FMT_INR2, bold=True)
dname("RO_LCOE", "Rollout Economics", f"E{rr}")
ws.freeze_panes = "A5"

# ================================================================ COMPARISON
ws = cmpw
title(ws, "How it compares on the same footprint", "Rollout area from 'Client'. A replaces a walkway and generates; D is the business-as-usual pair it should be compared with.")
widths(ws, {"A": 46, "B": 18, "C": 18, "D": 18, "E": 18, "F": 4, "G": 8, "H": 12, "I": 12})
header(ws, 4, ["Metric", "A. GridMats Solar Walkway", "B. Rooftop solar", "C. Solar carport", "D. Paver walkway + rooftop"])
# PV factors (years 1-20) for degrading, escalating savings, and escalating O&M
put(ws, "G4", "Year", bold=True); put(ws, "H4", "PVF walkway", bold=True); put(ws, "I4", "PVF rooftop", bold=True)
for y in range(1, 21):
    rr = 4 + y
    put(ws, f"G{rr}", y, FMT_NUM)
    put(ws, f"H{rr}", f"=(1-Deg1)*(1-Degr)^(G{rr}-1)*(1+TariffEsc)^(G{rr}-1)/(1+DiscRate)^G{rr}", '0.0000')
    put(ws, f"I{rr}", f"=(1-Deg1)*(1-RoofDegr)^(G{rr}-1)*(1+TariffEsc)^(G{rr}-1)/(1+DiscRate)^G{rr}", '0.0000')
put(ws, "G25", "Σ", bold=True)
put(ws, "H25", "=SUM(H5:H24)", '0.000'); put(ws, "I25", "=SUM(I5:I24)", '0.000')
put(ws, "G26", "Annuity", bold=True)
put(ws, "H26", "=(1-(1+DiscRate)^-20)/DiscRate", '0.000')
dname("PVF_W", "Comparison", "H25"); dname("PVF_R", "Comparison", "I25"); dname("ANN20", "Comparison", "H26")
rows = [
    ("Footprint (m²)", "=CampusArea", "=CampusArea", "=CampusArea", "=CampusArea", FMT_NUM),
    ("Installed capacity (kWp)", "=RO_KWp", "=CampusArea*RoofDens/1000", "=CampusArea*RoofDens/1000", "=CampusArea*RoofDens/1000", FMT_NUM1),
    ("Total capex (₹)", "=RO_Capex", "=C7*1000*RoofWp", "=D7*1000*CarportWp", "=CampusArea*ConvWalk+C7*1000*RoofWp", FMT_INR),
    ("Capex per kWp (₹)", "=B8/B7", "=C8/C7", "=D8/D7", "=E8/E7", FMT_INR),
    ("Year-1 generation (kWh)", "=RO_Gen1", "=C7*SunYear*RoofPR", "=D7*SunYear*RoofPR", "=E7*SunYear*RoofPR", FMT_NUM),
    ("Year-1 savings (₹)", "=B10*Tariff", "=C10*Tariff", "=D10*Tariff", "=E10*Tariff", FMT_INR),
    ("Year-1 O&M (₹)", "=OMPct*B8", "=RoofOM*C8", "=CarportOM*D8", "=RoofOM*C8+ConvWalkOM*CampusArea*ConvWalk", FMT_INR),
    ("Simple payback on year-1 figures (years)", "=B8/(B11-B12)", "=C8/(C11-C12)", "=D8/(D11-D12)", "=E8/(E11-E12)", FMT_YRS),
    ("NPV over 20 years (₹)", "=-B8+B11*PVF_W-B12*ANN20", "=-C8+C11*PVF_R-C12*ANN20", "=-D8+D11*PVF_R-D12*ANN20", "=-E8+E11*PVF_R-E12*ANN20", FMT_INR),
    ("Provides a walking surface", "Yes", "No", "Shade, not a walkway", "Yes", None),
    ("Needs roof or extra structure", "No", "Roof", "Steel canopy and footings", "Roof", None),
    ("Recycled plastic in the structure (t)", "=CampusArea*PlasticLW/1000", 0, 0, 0, FMT_NUM1),
]
for i, (lab, *vals, fmt) in enumerate(rows):
    r = 6 + i
    put(ws, f"A{r}", lab, kind="label", bold=i in (7, 8))
    for c, v in zip("BCDE", vals):
        put(ws, f"{c}{r}", v, fmt, bold=i in (7, 8), align="center" if isinstance(v, str) and not v.startswith("=") else None)
for c, key in zip("BCDE", ["W", "Roof", "Carport", "ConvRoof"]):
    dname(f"CMP_PerKWp_{key}", "Comparison", f"{c}9")
    dname(f"CMP_Payback_{key}", "Comparison", f"{c}13")
    dname(f"CMP_NPV_{key}", "Comparison", f"{c}14")
    dname(f"CMP_Capex_{key}", "Comparison", f"{c}8")
note(ws, "A19", "Rooftop and carport figures are the earlier ROI model's benchmarks. Rooftop is cheaper per kWp; the walkway earns its place where roofs are full, shaded or carry plant, and where a walkway is being built or relaid anyway (compare A with D).")
ws.merge_cells("A19:E21")

# ================================================================ SENSITIVITY
ws = sens
title(ws, "Sensitivity: rollout payback (years)", "Simple payback on year-1 figures for the campus rollout. Each table varies two inputs and keeps the rest.")
widths(ws, {"A": 26})
for i in range(2, 10):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 11
fill_good = PatternFill("solid", fgColor="E3F1E6")
fill_mid = PatternFill("solid", fgColor="FFF4D6")
fill_bad = PatternFill("solid", fgColor="FBEAEC")


def grid(r0, title_txt, row_label, row_vals, row_fmt, col_label, col_vals, col_fmt, formula, name_prefix):
    section(ws, r0, title_txt, ncols=1 + len(col_vals))
    put(ws, f"A{r0+1}", f"{row_label} ↓  /  {col_label} →", kind="label", bold=True)
    for j, cv in enumerate(col_vals):
        c = openpyxl.utils.get_column_letter(2 + j)
        put(ws, f"{c}{r0+1}", cv, col_fmt, kind="input", bold=True, align="center")
    for i, rv in enumerate(row_vals):
        r = r0 + 2 + i
        put(ws, f"A{r}", rv, row_fmt, kind="input", bold=True, align="center")
        for j in range(len(col_vals)):
            c = openpyxl.utils.get_column_letter(2 + j)
            f = formula.format(row=f"$A{r}", col=f"{c}${r0+1}")
            put(ws, f"{c}{r}", f"=IFERROR({f},\"n/a\")", '0.0', align="center")
    return r0 + 2, r0 + 1 + len(row_vals)


net_saving = "(PowerDens/1000*SunYear*{pr}*(1-ShadeLoss)*{tar}-OMPct*RO_CostM2)"
g1 = grid(4, "Power price × performance ratio", "₹ / kWh", [8, 9, 10, 11, 12, 13, 14], FMT_INR2, "Performance ratio",
          [0.60, 0.65, 0.70, 0.75, 0.80], FMT_PCT, "RO_NetM2/" + net_saving.format(pr="{col}", tar="{row}"), "S1")
g2 = grid(15, "Power price × change in the rollout PV-layer price", "₹ / kWh", [8, 9, 10, 11, 12, 13, 14], FMT_INR2, "PV-layer price change",
          [-0.4, -0.3, -0.2, -0.1, 0.0], '0%',
          "(RO_NetM2+RO_PV*{col}*(1+LogPct))/(PowerDens/1000*SunYear*PR*(1-ShadeLoss)*{row}-OMPct*(RO_CostM2+RO_PV*{col}*(1+LogPct)))", "S2")
section(ws, 26, "Pilot payback on the client's outlay × GridMats co-funding", ncols=6)
put(ws, "A27", "Co-funding of the PV layer →", kind="label", bold=True)
cof = [0, 0.25, 0.5, 0.75]
for j, v in enumerate(cof):
    c = openpyxl.utils.get_column_letter(2 + j)
    put(ws, f"{c}27", v, FMT_PCT, kind="input", bold=True, align="center")
    put(ws, f"{c}28", f"=IFERROR((PilotPays+(CoFund-{c}27)*PVPilot*PilotArea)/(PE_Save1-PE_OM1),\"n/a\")", '0.0', align="center")
put(ws, "A28", "Simple payback (years)", kind="label")
note(ws, "A30", "Shading: green ≤ 6 years, amber 6-9 years, red above 9 years (applied in Excel with conditional formatting).")
from openpyxl.formatting.rule import CellIsRule
for rng in (f"B{g1[0]}:F{g1[1]}", f"B{g2[0]}:F{g2[1]}", "B28:E28"):
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThanOrEqual", formula=["6"], fill=fill_good))
    ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["6.0001", "9"], fill=fill_mid))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["9"], fill=fill_bad))

# ================================================================ SCALE-UP
ws = sc
title(ws, "Scale-up: from pilot to campus", "Year-1 figures. Pilot uses pilot prices; later phases use the rollout prices for their area.")
widths(ws, {"A": 46, "B": 16, "C": 16, "D": 16, "E": 16, "F": 40})
header(ws, 4, ["Measure", "Pilot", "Phase 2", "Campus", "Portfolio", "Note"])
put(ws, "A5", "Area (m²)", kind="label", bold=True)
for c, f in zip("BCDE", ["=PilotArea", "=Phase2Area", "=CampusArea", "=PortfolioArea"]):
    put(ws, f"{c}5", f, FMT_NUM, kind="link", bold=True)
_civ = 'IF(BuildType="New walkway",CivilNew,CivilOverlay)'
_grd = 'IF({c}5<1000,GridT1,IF({c}5<2500,GridT2,IF({c}5<5000,GridT3,GridT4)))'
_pv = 'IF({c}5>=ScaleFrom,PVScale,PVPilot)'
_bos = 'IF({c}5>=ScaleFrom,BOSScale,BOSPilot)*PowerDens'
_inst = 'IF({c}5>=ScaleFrom,InstScale,InstPilot)'
SC_COST = ("=" + _civ + "+" + _grd + "+" + _pv + "+" + _bos + "+LogPct*(" + _civ + "+" + _grd + "+" + _pv + "+" + _bos + ")"
           "+ROUNDUP({c}6/ConnKWp,0)*IntegLump/{c}5+" + _inst)
srows = [
    ("Rated power (kWp)", "={c}5*PowerDens/1000", FMT_NUM1),
    ("Generation, year 1 (MWh)", "={c}6*SpecYield/1000", FMT_NUM1),
    ("Electricity saved, year 1 (₹ lakh)", "={c}7*1000*Tariff/100000", FMT_NUM2),
    ("CO2 avoided per year (t)", "={c}7*GridEF", FMT_NUM1),
    ("Installed cost per m² (₹)", SC_COST, FMT_INR),
    ("Installed cost (₹ lakh)", "={c}10*{c}5/100000", FMT_NUM2),
    ("Net of walkway credit (₹ lakh)", "=({c}10-IF(BuildType=\"New walkway\",ConvWalk,0))*{c}5/100000", FMT_NUM2),
    ("Simple payback on full cost (years)", "={c}12/({c}8-OMPct*{c}11)", FMT_YRS),
    ("Recycled PP in the base (t)", "={c}5*PlasticLW/1000", FMT_NUM2),
]
for i, (lab, f, fmt) in enumerate(srows):
    r = 6 + i
    put(ws, f"A{r}", lab, kind="label", bold=i == 8)
    for c in "BCDE":
        put(ws, f"{c}{r}", f.format(c=c), fmt, bold=i == 8)
for c, key in zip("BCDE", ["Pilot", "P2", "Campus", "Port"]):
    dname(f"SC_KWp_{key}", "Scale-up", f"{c}6")
    dname(f"SC_MWh_{key}", "Scale-up", f"{c}7")
    dname(f"SC_Save_{key}", "Scale-up", f"{c}8")
    dname(f"SC_CO2_{key}", "Scale-up", f"{c}9")
    dname(f"SC_Cost_{key}", "Scale-up", f"{c}11")
    dname(f"SC_Payback_{key}", "Scale-up", f"{c}13")
note(ws, "F13", "Pilot column shows full physical cost, before co-funding.")

# ================================================================ SCORECARD
ws = card
title(ws, "Pilot scorecard: measured, not assumed", "Agreed with the client before installation; filled in at the evaluation review.")
widths(ws, {"A": 40, "B": 34, "C": 40, "D": 16, "E": 10, "F": 10, "G": 30})
header(ws, 4, ["KPI", "Target / benchmark", "How it is measured", "Measured result", "Weight", "Score (1-5)", "Notes"])
kpis = [
    ("Energy generation vs estimate", "≥ 90% of the model's monthly estimate", "Energy meter and logger, monthly", 0.20),
    ("Performance ratio achieved", "≥ PR used in the model", "Generation ÷ (reference-cell irradiation × kWp)", 0.10),
    ("Pedestrian and maintenance-cart load", "4 kN/m² uniform and 4.5 kN point load (IS 875 Part 2); no cracking or movement", "Load test at handover; inspection", 0.15),
    ("Surface temperature (peak)", "Within module specification", "Module-temperature sensors", 0.05),
    ("Slip resistance and walking comfort", "R11 (DIN 51130) or wet pendulum PTV ≥ 36", "Pendulum test wet and dry; user feedback", 0.10),
    ("Drainage and water ingress", "No ponding, no ingress", "Inspection after heavy rain", 0.10),
    ("Electrical uptime", "≥ 98%", "Logger availability data", 0.10),
    ("Cleaning and maintenance effort", "Within O&M assumption", "Cleaning log", 0.05),
    ("Module survival and repairability", "No irreparable failures", "Replacement log", 0.10),
    ("Visual finish and user acceptance", "Client acceptance", "Facilities and user survey", 0.05),
]
for i, (k, tgt, how, wgt) in enumerate(kpis):
    r = 5 + i
    put(ws, f"A{r}", k, kind="label")
    put(ws, f"B{r}", tgt, kind="note")
    put(ws, f"C{r}", how, kind="note")
    put(ws, f"D{r}", None, kind="input")
    put(ws, f"E{r}", wgt, FMT_PCT, kind="input")
    put(ws, f"F{r}", None, kind="input")
put(ws, "A15", "Weight total (must be 100%)", kind="label", bold=True)
put(ws, "E15", "=SUM(E5:E14)", FMT_PCT, bold=True)
put(ws, "A16", "Weighted score (out of 5)", kind="label", bold=True)
put(ws, "F16", "=IFERROR(SUMPRODUCT(E5:E14,F5:F14)/SUM(E5:E14),\"-\")", '0.0', bold=True)
note(ws, "A18", "KPIs and weights from the earlier GridMats Solar Walkway ROI model (Reliance / Jio R&D pilot). Fill in columns D and F at the review.")
ws.merge_cells("A18:G19")

# ================================================================ SOURCES & NOTES
ws = src
title(ws, "Sources and notes", "Where every number comes from, and what changed from earlier GridMats calculations.")
widths(ws, {"A": 44, "B": 30, "C": 30, "D": 70})
section(ws, 4, "Changes from earlier GridMats Solar Walkway calculations", ncols=4)
header(ws, 5, ["Item", "Earlier value (file)", "This model", "Why"])
changes = [
    ("Rated power of the walkway", "126 W/m² (ROI model: 580 × 450 mm, 82% coverage); 110 W/m² (2025 decks, flier to 6 Oct)", "160 W/m²", "MD correction, 6 Oct 2026: 40 W module on each grid, 4 per m²."),
    ("PV layer price", "₹2,500/m² (ROI model input cell); ₹9,500/m² quote; ₹8,000-9,000/m² range (Jul 2026)", "Built up from actual unit costs", "GridMats' actual material cost is ₹3,075 per 40 W piece (₹12,300/m²), above the old ₹9,500/m² quote. Price now = cost + wiring, inverter, labour, management and margin ('Cost Build-up')."),
    ("Structural base", "GridMats HD ₹1,200/m² (ROI model)", "OC-LW 40 at ₹690/m² (under 1,000 m²)", "Published price ladder (25 Sep 2026); actual cost ₹125 per grid (₹500/m²). The module sits on the 500 × 500 mm grid."),
    ("Performance ratio", "78%", "About 73% base (seven-loss stack)", "Flat, walk-on modules run hotter and collect more dust than tilted rooftop panels."),
    ("GST", "13.8% composite", "8.9% composite (confirm)", "GST on solar goods cut from 12% to 5% from 22 Sep 2025."),
    ("Power price", "₹9/kWh placeholder", "Client sheet", "Use the landed rate from the site's electricity bill."),
    ("Cost calculator (Apr 2025)", "₹6,540-7,800/m² at 80-111 W/m²; ₹70-84/W", "Superseded", "Lower power density and older module; kept for reference."),
    ("Cell technology and losses", "Not modelled; flat 78% PR, 0.5%/yr", "TOPCon (n-type) library; PR built from seven losses", "Latest TOPCon data: −0.29%/°C, 1% then 0.4%/yr; heat and dust matter more on a flat walkway."),
]
for i, row in enumerate(changes):
    for c, v in zip("ABCD", row):
        put(ws, f"{c}{6+i}", v, kind="label" if c == "A" else "note")
        ws[f"{c}{6+i}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[6 + i].height = 42
r = 6 + len(changes) + 1
section(ws, r, "Sources", ncols=4)
sources = [
    ("GridMats Solar Walkway flier (6 Oct 2026)", "Design basis: 40 W per module, 4 per m², ~160 W/m², 15-16% surface efficiency; pilot method; status."),
    ("GridMats Solar Walkway Client/Partner Deck v2", "Construction (PV layer, 40 × 40 mm MLP lumber, GridMats base), 550 × 400 mm module, pilot pathway."),
    ("GridMats Solar Walkway ROI Comparison Model (Reliance / Jio R&D pilot)", "Cost stack, rooftop and carport benchmarks, O&M, pilot objectives and scorecard."),
    ("GridMats Solar cost calculator (Apr 2025)", "Earlier module, MLP and labour costs; superseded."),
    ("GridMats actual unit costs per 40 W piece (7 Oct 2026)", "Module ₹1,200, glass ₹1,600, GridMats grid ₹125, MLP lumber ₹150. Excludes margin, labour, wiring, inverter and management."),
    ("GridMats price ladder (25 Sep 2026); Design Engine 73 unit rates", "OC-LW 40 prices; base-works rates."),
    ("GridMats financials (Jul 2026)", "Solar walkway pricing range ₹8,000-9,000/m²."),
    ("CEA CO2 Baseline Database v21.0", "Grid emission factor 0.710 t CO2/MWh."),
    ("Income-tax Act; GST Council (Sep 2025), Notification 9/2025-CT(R); Gujarat AAR GUJ/GAAR/R/2024/01", "40% WDV depreciation; 5% GST on solar goods; ITC on captive solar."),
    ("pv-magazine (25 Feb 2026); taiyangnews; JMK Research (Apr 2026); manufacturer datasheets", "TOPCon cell and module efficiency, temperature coefficient, degradation and price."),
    ("MERC order, Case 75 of 2025 (25 Mar 2026); MSEDCL circulars", "HT tariffs, solar-hour ToD rebate, electricity duty."),
    ("NASA-derived Navi Mumbai climatology; Meteonorm 8.2 (ICADMSES 2026)", "Mumbai sunshine calibration."),
    ("IS 875 (Part 2); DIN 51130; BS 7976-2; IS 14286 and IS/IEC 61730 (BIS)", "Load, slip-resistance and module-registration targets."),
]
for i, (s, u) in enumerate(sources):
    put(ws, f"A{r+1+i}", s, kind="label")
    put(ws, f"D{r+1+i}", u, kind="note")
note(ws, f"A{r+len(sources)+2}", "Status: the GridMats Solar Walkway is pre-commercial; pilot validation is under way and no commercial deployment is claimed. Yield, life and upkeep are estimates the pilot measures.")
ws.merge_cells(f"A{r+len(sources)+2}:D{r+len(sources)+3}")

# ================================================================ COVER
ws = cover
ws.sheet_view.showGridLines = False
widths(ws, {"A": 3, "B": 60, "C": 22, "D": 4, "E": 40})
ws["B2"] = "GridMats® Solar Walkway"
ws["B2"].font = Font(name=FONT, size=22, bold=True, color="C22935")
ws["B3"] = "Pilot and rollout cost-benefit model"
ws["B3"].font = Font(name=FONT, size=13, color="3F3F3F")
put(ws, "B5", "Prepared for", kind="label", bold=True); put(ws, "C5", "=ClientEntity", kind="link", bold=True)
put(ws, "B6", "Site", kind="label"); put(ws, "C6", "=SiteName&\", \"&SiteCity", kind="link")
put(ws, "B7", "Date / prepared by", kind="label"); put(ws, "C7", "=PropDate&\" · \"&PreparedBy", kind="link")
put(ws, "B8", "Scenario", kind="label"); put(ws, "C8", "=Scenario", kind="link")
section(ws, 10, "Headline results", ncols=4, col=2)
cov = [
    ("Pilot", None, None),
    ("Pilot area (m²) / rated power (kWp)", "=TEXT(PilotArea,\"0\")&\" m² · \"&TEXT(PilotKWp,\"0.0\")&\" kWp\"", None),
    ("Generation in year 1 (kWh)", "=PilotKWh1", FMT_NUM),
    ("Client pays, ex-GST (₹)", "=PilotPays", FMT_INR),
    ("GridMats contribution (₹)", "=PilotContrib", FMT_INR),
    ("Pilot payback on the client's outlay (years)", "=PE_Payback", FMT_YRS),
    ("Campus rollout", None, None),
    ("Rollout area (m²)", "=CampusArea", FMT_NUM),
    ("Installed cost (₹/sq ft)", "=RO_CostSqft", FMT_INR),
    ("Simple payback (years)", "=RO_Payback", FMT_YRS),
    ("IRR over 20 years", "=RO_IRR", FMT_PCT),
    ("Comparison (same footprint)", None, None),
    ("Payback: walkway / rooftop / paver walkway + rooftop (years)",
     "=TEXT(CMP_Payback_W,\"0.0\")&\" / \"&TEXT(CMP_Payback_Roof,\"0.0\")&\" / \"&TEXT(CMP_Payback_ConvRoof,\"0.0\")", None),
]
r = 11
for lab, f, fmt in cov:
    if f is None:
        put(ws, f"B{r}", lab, kind="label", bold=True)
        for c in "BCDE":
            ws[f"{c}{r}"].fill = PatternFill("solid", fgColor="F6F6F4")
    else:
        put(ws, f"B{r}", lab, kind="label")
        put(ws, f"C{r}", f, fmt, kind="link", bold=True)
    r += 1
r += 1
section(ws, r, "How to use this template", ncols=4, col=2)
how = [
    "For a new client, change the yellow cells on 'Client' (or build from clients/<id>.json with build/build.sh). Nothing else needs editing.",
    "GridMats product prices and assumptions sit on 'Product & Costs'. Unit costs, margin and the cost-plus price build-up sit on 'Cost Build-up': internal, remove before sending to a client.",
    "Pick the sunshine profile on 'Client'; replace the city row on 'Irradiance' with site data after the site assessment.",
    "Blue = input · yellow fill = key assumption to confirm · black = formula · green = link to another sheet.",
    "All money in ₹, ex-GST unless stated. The Solar Walkway is pre-commercial: yield, life and upkeep are estimates the pilot measures.",
]
for i, h in enumerate(how):
    put(ws, f"B{r+1+i}", f"{i+1}. {h}", kind="label")
    ws[f"B{r+1+i}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(f"B{r+1+i}:E{r+1+i}")
    ws.row_dimensions[r + 1 + i].height = 30

for w in wb.worksheets:
    w.sheet_view.zoomScale = 90
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
