"""Pull chart and table series out of a recalculated Solar Walkway workbook."""
import json
import sys

import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
d = {}

ws = wb["Energy"]
d["months"] = [ws[f"A{r}"].value for r in range(11, 23)]
d["monthly_kwh"] = [ws[f"E{r}"].value for r in range(11, 23)]
d["monthly_ghi"] = [ws[f"C{r}"].value for r in range(11, 23)]

ws = wb["Rollout Economics"]
for r in range(20, 60):
    if ws[f"A{r}"].value == 0 and ws[f"H{r}"].value is not None:
        d["rollout_cum"] = [ws[f"I{r+y}"].value for y in range(21)]
        break

ws = wb["Pilot BoQ"]
d["boq"] = []
for r in range(5, 17):
    if ws[f"A{r}"].value:
        d["boq"].append({"item": ws[f"A{r}"].value, "list": ws[f"E{r}"].value, "contrib": ws[f"F{r}"].value,
                         "pays": ws[f"G{r}"].value, "note": ws[f"H{r}"].value})

ws = wb["Sensitivity"]
d["sens_pr"] = {"cols": [ws.cell(row=5, column=c).value for c in range(2, 7)],
                "rows": [[ws.cell(row=r, column=c).value for c in range(1, 7)] for r in range(6, 13)]}
d["sens_price"] = {"cols": [ws.cell(row=16, column=c).value for c in range(2, 6)],
                   "rows": [[ws.cell(row=r, column=c).value for c in range(1, 6)] for r in range(17, 24)]}
d["sens_cofund"] = {"cols": [ws.cell(row=27, column=c).value for c in range(2, 6)],
                    "vals": [ws.cell(row=28, column=c).value for c in range(2, 6)]}

ws = wb["Comparison"]
d["comparison"] = {ws[f"A{r}"].value: [ws[f"{c}{r}"].value for c in "BCDE"] for r in range(6, 18)}

ws = wb["Scale-up"]
d["scale"] = {ws[f"A{r}"].value: [ws[f"{c}{r}"].value for c in "BCDE"] for r in range(5, 15)}

ws = wb["Scorecard"]
d["kpis"] = [{"kpi": ws[f"A{r}"].value, "target": ws[f"B{r}"].value, "how": ws[f"C{r}"].value, "weight": ws[f"E{r}"].value}
             for r in range(5, 15)]
json.dump(d, open(sys.argv[2], "w"), indent=1, default=str)
print("chartdata ok:", ", ".join(d))
