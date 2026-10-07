"""Styling helpers for the GridMats x Reliance cost-benefit workbook."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

FONT = "Arial"
GREEN_DARK = "1A1A1A"  # brand 90% black (title/header text)
GREEN = "3F3F3F"  # brand 75% black (section bars)
GREEN_LIGHT = "FBEAEC"  # brand red-soft (table headers)
RED = "C22935"  # GridMats Brick Red
SUN = "F2A900"
SUN_LIGHT = "FFF4D6"
GREY = "5B6770"
GREY_LIGHT = "F3F5F6"
WHITE = "FFFFFF"

thin = Side(style="thin", color="D0D7DC")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=Side(style="thin", color="9AA5AD"))
TOP2 = Border(top=Side(style="thin", color="1F2A30"), bottom=Side(style="double", color="1F2A30"))

INPUT_FONT = Font(name=FONT, size=10, color="0000FF")
LINK_FONT = Font(name=FONT, size=10, color="008000")
CALC_FONT = Font(name=FONT, size=10, color="000000")
BOLD = Font(name=FONT, size=10, bold=True)
NOTE_FONT = Font(name=FONT, size=9, italic=True, color=GREY)
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
INPUT_FILL = PatternFill("solid", fgColor="FFFDE7")

FMT_INR = '₹#,##0;(₹#,##0);"-"'
FMT_INR2 = '₹#,##0.00;(₹#,##0.00);"-"'
FMT_LAKH = '₹#,##0.00" L";(₹#,##0.00" L");"-"'
FMT_CR = '₹#,##0.00" Cr";(₹#,##0.00" Cr");"-"'
FMT_NUM = '#,##0;(#,##0);"-"'
FMT_NUM1 = '#,##0.0;(#,##0.0);"-"'
FMT_NUM2 = '#,##0.00;(#,##0.00);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_X = '0.0"x";(0.0"x");"-"'
FMT_YRS = '0.0" yrs";(0.0" yrs");"-"'


def title(ws, text, sub=None, width_cols=10):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=RED)
    ws.row_dimensions[1].height = 26
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FONT, size=10, italic=True, color=GREY)
    ws.sheet_view.showGridLines = False


def section(ws, row, text, ncols=8, col=1):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(name=FONT, size=11, bold=True, color=WHITE)
    for j in range(col, col + ncols):
        ws.cell(row=row, column=j).fill = PatternFill("solid", fgColor=GREEN)
    ws.row_dimensions[row].height = 18


def header(ws, row, labels, col=1, fill=GREEN_LIGHT, wrap=True):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=lab)
        c.font = Font(name=FONT, size=10, bold=True, color=GREEN_DARK)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center" if j else "left", vertical="center", wrap_text=wrap)
        c.border = BOTTOM


def put(ws, ref, value, fmt=None, kind="calc", bold=False, comment=None, key=False, align=None):
    """kind: input | calc | link | label | note"""
    c = ws[ref]
    c.value = value
    if kind == "input":
        c.font = Font(name=FONT, size=10, color="0000FF", bold=bold)
        c.fill = KEY_FILL if key else INPUT_FILL
        c.border = BOX
    elif kind == "link":
        c.font = Font(name=FONT, size=10, color="008000", bold=bold)
    elif kind == "note":
        c.font = NOTE_FONT
    elif kind == "label":
        c.font = Font(name=FONT, size=10, bold=bold)
    else:
        c.font = Font(name=FONT, size=10, bold=bold)
    if fmt:
        c.number_format = fmt
    if comment:
        c.comment = Comment(comment, "GridMats model")
        c.comment.width = 320
        c.comment.height = 120
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    return c


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w
