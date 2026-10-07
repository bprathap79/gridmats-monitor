"""Build the GridMats Solar x Reliance cost-benefit workbook (formula-driven)."""
import sys
import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment
from xlhelpers import *

OUT = sys.argv[1] if len(sys.argv) > 1 else "model.xlsx"
wb = openpyxl.Workbook()
N = {}  # defined name -> absolute ref


def dname(name, sheet, ref):
    s = f"'{sheet}'" if " " in sheet or "-" in sheet else sheet
    col = "".join(ch for ch in ref if ch.isalpha())
    row = "".join(ch for ch in ref if ch.isdigit())
    attr = f"{s}!${col}${row}"
    wb.defined_names[name] = DefinedName(name, attr_text=attr)
    N[name] = attr


YEARS = list(range(26))          # 0..25
YC = {t: L(3 + t) for t in YEARS}  # year t -> column letter (C..AB)
FIRST, Y1, LAST, PVCOL = YC[0], YC[1], YC[25], L(29)  # C, D, AB, AC

# ---------------------------------------------------------------- sheets
cover = wb.active
cover.title = "Cover"
inp = wb.create_sheet("Inputs")
ur = wb.create_sheet("Unit Rates")
lcc = wb.create_sheet("Road LCC")
park = wb.create_sheet("Solar Park 100MW")
walk = wb.create_sheet("Solar Walkway")
sus = wb.create_sheet("Sustainability")
pilot = wb.create_sheet("Pilot Offer")
scale = wb.create_sheet("Scale-up")
sens = wb.create_sheet("Sensitivity")
src = wb.create_sheet("Sources")

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = "C22935" if ws.title in ("Cover", "Pilot Offer") else (
        "FFC000" if ws.title == "Inputs" else "3F3F3F")

# ================================================================ INPUTS
ws = inp
title(ws, "Inputs and assumptions",
      "Blue = input you can change · yellow fill = key assumption · black = formula. All money ex-GST, 2026 prices.")
widths(ws, {"A": 58, "B": 15, "C": 14, "D": 13, "E": 13, "F": 13, "G": 100})
put(ws, "A3", "Scenario used across the model", kind="label", bold=True)
put(ws, "C3", "Base", kind="input", key=True, bold=True, align="center")
put(ws, "D3", "← choose Conservative, Base or Optimistic", kind="note")
dv = DataValidation(type="list", formula1='"Conservative,Base,Optimistic"', allow_blank=False)
ws.add_data_validation(dv); dv.add("C3")
put(ws, "A4", "Scenario index", kind="label")
header(ws, 6, ["Assumption", "Unit", "Value used", "Conservative", "Base", "Optimistic", "Source / basis"])
put(ws, "C4", "=MATCH(C3,$D$6:$F$6,0)", FMT_NUM)
dname("ScnIdx", "Inputs", "C4")
dname("Scenario", "Inputs", "C3")

row = 7
SRC_NOTES = []


def sec(text):
    global row
    row += 1
    section(ws, row, text, ncols=7)
    row += 1


def inp_row(name, label, unit, val, fmt, source, key=False, dv_list=None):
    """val: number, str, or (cons, base, opt) tuple."""
    global row
    put(ws, f"A{row}", label, kind="label")
    put(ws, f"B{row}", unit, kind="note")
    if isinstance(val, tuple):
        for col, v in zip("DEF", val):
            put(ws, f"{col}{row}", v, fmt, kind="input", key=key)
        put(ws, f"C{row}", f"=INDEX(D{row}:F{row},ScnIdx)", fmt, bold=True)
    else:
        put(ws, f"C{row}", val, fmt, kind="input", key=key, align="center" if isinstance(val, str) else None)
        if dv_list:
            d = DataValidation(type="list", formula1=f'"{dv_list}"', allow_blank=False)
            ws.add_data_validation(d); d.add(f"C{row}")
    put(ws, f"G{row}", source, kind="note")
    ws[f"G{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    dname(name, "Inputs", f"C{row}")
    row += 1


sec("1. General")
inp_row("DiscRate", "Discount rate (nominal, pre-tax)", "%", 0.10, FMT_PCT,
        "Assumption: indicative large-corporate cost of capital. Sensitivity sheet shows 8% and 12%.", key=True)
inp_row("CostEsc", "Escalation of civil and upkeep costs", "% / yr", 0.05, FMT_PCT, "Assumption (close to Indian CPI).")
inp_row("TaxRate", "Corporate income-tax rate", "%", 0.2517, FMT_PCT,
        "Section 115BAA rate incl. surcharge and cess. Used only in the walkway after-tax view.")
inp_row("CarbonPrice", "Carbon value", "₹ / t CO2e", (400, 700, 1500), FMT_INR,
        "India CCTS: first carbon credit certificate trades expected around Oct 2026; expected ₹250-1,500/t, scenario studies ₹400-900/t (reclimatize.in, 2026).")
inp_row("CarbonEsc", "Carbon value escalation", "% / yr", 0.05, FMT_PCT, "Assumption.")
inp_row("GridEF", "India grid emission factor", "t CO2 / MWh", 0.710, FMT_NUM2,
        "CEA CO2 Baseline Database v21.0 (FY2024-25), weighted average.")

sec("2. Solar park block (Kutch region)")
inp_row("BlockMWp", "Solar block size used for comparison", "MWp", 100, FMT_NUM, "Basis of comparison.")
inp_row("SolarCapex", "Solar plant capex (for context only)", "₹ Cr / MW", 4.0, FMT_NUM2,
        "Utility-scale benchmark ₹3.5-4.5 Cr/MW (greentechlead, 2026).")
inp_row("Yield", "Specific yield, Kutch region", "kWh / kWp / yr", 1870, FMT_NUM,
        "Gujarat plant 1,872 kWh/kWp at 21.3% CUF; Kutch plane-of-array irradiance 2,100-2,350 kWh/m²/yr.")
inp_row("Degr", "Module degradation", "% / yr", 0.004, FMT_PCT,
        "TOPCon datasheets 0.4%/yr after year 1; Reliance HJT 0.3%/yr (company claim).")
inp_row("EnergyValue", "Value of each extra kWh generated", "₹ / kWh", (2.55, 2.70, 2.90), FMT_INR2,
        "Solar PPA ₹2.52-2.61 (NHPC/NTPC); solar + storage ₹2.86-2.87 (SECI ISTS-XX, Oct 2025). Conservative: captive value to Reliance is higher.", key=True)
inp_row("EnergyEsc", "Energy value escalation", "% / yr", 0.0, FMT_PCT, "Flat, PPA-style.")
inp_row("RoadKm", "Internal road per 100 MW", "km", 10, FMT_NUM1,
        "GridMats solar and wind parks flier: 8-12 km per 100 MW. Independent estimate 0.1-0.2 km/MW.", key=True)
inp_row("MainShare", "Share of road length that is main / spine / inverter-yard road", "%", 0.30, FMT_PCT, "Assumption.")
inp_row("MainWidth", "Main road width", "m", 6.0, FMT_NUM1, "GridMats per-km park rate covers 5-7 m roads; Pavagada DPR uses 8.5 m and 4.5 m.")
inp_row("TrackWidth", "O&M track width", "m", 4.5, FMT_NUM1, "Pavagada DPR 4.5 m; NTPC tenders 3.75 m WBM.")
inp_row("MainBaseline", "Conventional main-road surface it replaces", "choice", "Bituminous", None,
        "Choose Bituminous (default, conservative) or PQC. O&M tracks are compared with WBM / gravel.", dv_list="Bituminous,PQC")
inp_row("SoilingLoss", "Average yearly energy lost to module soiling, with routine cleaning", "% of yield", 0.03, FMT_PCT,
        "West-India dry-season soiling 0.3-0.5%/day (Gandhinagar 0.37%/day); pv-magazine India uses a 3% example (Apr 2026). Assumption.")
inp_row("DustShare", "Share of that soiling caused by traffic dust from unpaved O&M tracks", "%", (0.10, 0.15, 0.25), FMT_PCT,
        "Assumption. Univ. of Johannesburg (2025) measured PV soiling highest next to a gravel road, falling with distance. The pilot measures this.", key=True)
inp_row("DustCutS2", "Dust cut by GridMats S2 (graded chips confined in cells) vs WBM", "%", 0.70, FMT_PCT,
        "Assumption: 8-12 mm graded chips, no fines, held in cells. Pilot KPI.")
inp_row("CleanCost", "Module cleaning cost", "₹ lakh / MW / yr", 1.5, FMT_NUM2,
        "Robotic about ₹1 lakh/MW/yr; manual up to ₹7.5 lakh/MW/yr (Renewable Watch, 2023).")
inp_row("CleanSave", "Cleaning cost avoided through fewer cycles", "%", (0.0, 0.0, 0.05), FMT_PCT,
        "Optimistic case only, so it is not double-counted with the energy gain.")
inp_row("DaysEarly", "Earlier energisation from all-weather access", "days", (0, 0, 2), FMT_NUM, "Optimistic case only.")
inp_row("DustYears", "Years of heavy construction traffic needing dust suppression", "years", 2, FMT_NUM, "Assumption.")
inp_row("DustCost", "Dust suppression on WBM (water tanker + water)", "₹ / m² / yr", 75, FMT_INR,
        "Assumption: about ₹4,000 a day per tanker covering ~13,500 m², 250 days a year.")
inp_row("DustWater", "Sprinkling water on WBM", "L / m² / day", 1.0, FMT_NUM1, "Assumption.")
inp_row("DustDays", "Sprinkling days per year", "days", 200, FMT_NUM, "Assumption.")
inp_row("Rainfall", "Annual rainfall, Kutch / Patan", "mm", 400, FMT_NUM, "Assumption: Kutch district normal about 350-450 mm.")
inp_row("Infiltr", "Share of rain returned to ground on permeable S2", "%", 0.90, FMT_PCT, "GridMats Design Engine factor.")

sec("3. Conventional surfaces, build cost (₹ / m², ex-GST)")
inp_row("WBM_Capex", "WBM / gravel road (GSB 150 + WBM 150)", "₹ / m²", 600, FMT_INR,
        "Tamil Nadu G.O. (2021-22) single-layer WBM ₹21.99 lakh/km, about ₹590/m²; GridMats sales pitch ₹250-500/m² (surface only).")
inp_row("BT_Capex", "Bituminous road, IRC:37", "₹ / m²", 2000, FMT_INR,
        "CPWD DSR 2023 internal road (GSB 225 + WBM 200 + BM 50 + BC 40) about ₹2,050/m²; GridMats pitch ₹1,500-1,900/m².")
inp_row("PQC_Capex", "Concrete road, IRC:58 (PQC + DLC)", "₹ / m²", 3200, FMT_INR,
        "GridMats pitch ₹2,900-3,200/m²; Design Engine 73 (280 mm PQC) ₹3,561/m²; CPWD-derived ₹4,020/m².")
inp_row("Paver_Capex", "Paver-block walkway (avoided by a solar walkway)", "₹ / m²", 1100, FMT_INR,
        "Paver blocks 60-80 mm ₹900-1,300/m² (GridMats pitch); ₹548-1,374/m² (Infralens, 2026).")

sec("4. Upkeep (today's ₹ / m², escalated in the model)")
inp_row("WBM_Routine", "WBM: post-monsoon regrading, rut and pothole fill", "₹ / m² / yr", 40, FMT_INR,
        "Assumption (about ₹1.8 lakh/km/yr on a 4.5 m track). GridMats pitch: 'regrade every monsoon'.", key=True)
inp_row("WBM_Regravel", "WBM: regravelling", "₹ / m² / event", 250, FMT_INR, "Assumption, about 40% of build cost.")
inp_row("WBM_Cycle", "WBM: regravelling interval", "years", 4, FMT_NUM, "Regravelling every 4-7 years (TRL / World Bank road-cost data).")
inp_row("BT_Routine", "Bituminous: routine maintenance", "₹ / m² / yr", 20, FMT_INR, "IRC:82 routine maintenance ₹0.5-1.5 lakh/km/yr.")
inp_row("BT_Overlay", "Bituminous: 40 mm BC renewal", "₹ / m² / event", 450, FMT_INR, "Assumption.")
inp_row("BT_Cycle", "Bituminous: renewal interval", "years", 5, FMT_NUM, "IRC:82 renewal every 4-6 years; GridMats pitch 3-5 years.")
inp_row("PQC_Routine", "Concrete: routine maintenance", "₹ / m² / yr", 10, FMT_INR, "Assumption.")
inp_row("PQC_Repair", "Concrete: joint reseal, grinding, slab repairs", "₹ / m² / event", 150, FMT_INR, "Design Engine 73 lifecycle (repairs at years 10, 20, 30).")
inp_row("PQC_Cycle", "Concrete: repair interval", "years", 10, FMT_NUM, "Design Engine 73 lifecycle.")
inp_row("GM_RoutinePct", "GridMats: routine upkeep", "% of build cost / yr", 0.005, FMT_PCT, "Design Engine 73 lifecycle method (0.5% a year).")
inp_row("S1_TopUp", "GridMats S1: chip-and-cement top-up", "₹ / m² / event", 150, FMT_INR, "GridMats SOP: top-up at years 10, 20, 30.")
inp_row("S1_Cycle", "GridMats S1: top-up interval", "years", 10, FMT_NUM, "GridMats SOP.")
inp_row("S2_TopUp", "GridMats S2: chip top-up", "₹ / m² / event", 35, FMT_INR, "Assumption: about 10 L/m² of chips plus labour.")
inp_row("S2_Cycle", "GridMats S2: chip top-up interval", "years", 3, FMT_NUM, "Assumption (top-up visit at 6 months is included in the build).")

sec("5. GridMats rates and factors")
inp_row("GridHD", "GridMats OC-HD 40, park / annual-contract rate", "₹ / m²", 900, FMT_INR,
        "GridMats price ladder (25 Sep 2026): ₹900/m² above 25,000 m² or on a 12-month rate contract; ₹935 at 5,000-25,000 m².", key=True)
inp_row("GridLW", "GridMats OC-LW 40, annual rate", "₹ / m²", 550, FMT_INR, "GridMats price ladder (25 Sep 2026).")
inp_row("GridLWSmall", "GridMats OC-LW 40, under 1,000 m²", "₹ / m²", 690, FMT_INR, "GridMats price ladder (25 Sep 2026).")
inp_row("Geotex", "Geotextile, 200 gsm non-woven", "₹ / m²", 55, FMT_INR, "As supplied to Reliance in 2025.")
inp_row("Freight", "Freight, factory to Kutch site", "₹ / m²", 60, FMT_INR, "Design Engine BoQ indicative; at cost.")
inp_row("LayLab", "Grid laying labour", "₹ / m²", 60, FMT_INR, "Design Engine 73 BoQ.")
inp_row("NailPrice", "U-nail (8 mm bar)", "₹ each", 8, FMT_INR, "Design Engine 73 BoQ.")
inp_row("NailDensity", "U-nails per m² on roads", "nos / m²", 0.4, FMT_NUM2, "GridMats SOP: one per 2-3 m² on roads and slopes.")
inp_row("SubgradeRate", "Subgrade preparation", "₹ / m²", 60, FMT_INR, "Design Engine 73 BoQ (Bengaluru 2026 rates).")
inp_row("GSBRate", "GSB (MoRTH 401)", "₹ / m³", 1400, FMT_INR, "Design Engine 73 BoQ.")
inp_row("OGSBRate", "Open-graded GSB", "₹ / m³", 1600, FMT_INR, "Design Engine 73 BoQ.")
inp_row("WMMRate", "WMM (MoRTH 406)", "₹ / m³", 1900, FMT_INR, "Design Engine 73 BoQ.")
inp_row("SandRate", "M-sand / crusher-dust bedding", "₹ / m³", 1800, FMT_INR, "Design Engine 73 BoQ.")
inp_row("ChipRate", "6-12 mm stone chips (loose)", "₹ / m³", 1900, FMT_INR, "GridMats sample proposal (Design Engine 73).")
inp_row("CTGInfill", "S1 cement-treated gravel infill (chips, cement, mix, place, cure)", "₹ / m²", 253, FMT_INR,
        "Design Engine 73 BoQ, internal road: ₹7.58 lakh for 3,000 m².")
inp_row("S1Edges", "S1 edges and drains", "₹ / m²", 28, FMT_INR, "Design Engine 73 BoQ, internal road.")
inp_row("S2Edges", "S2 / walkway edges", "₹ / m²", 12, FMT_INR, "GridMats sample proposal.")
inp_row("PlasticHD", "Recycled PP in OC-HD 40", "kg / m²", 6.6, FMT_NUM1, "GridMats master data: 1.7 kg per panel, 3.9 panels/m².")
inp_row("PlasticLW", "Recycled PP in OC-LW 40", "kg / m²", 4.0, FMT_NUM1, "GridMats master data.")
inp_row("PPCredit", "CO2e saved per kg of recycled PP used", "kg CO2e / kg", 1.2, FMT_NUM1, "GridMats Design Engine factor.")
inp_row("CementEF", "Cement emission factor", "t CO2e / t", 0.65, FMT_NUM2, "GridMats Design Engine; India average 0.66-0.68.")
inp_row("PQC_Cement", "Cement in a PQC 220 + DLC 150 road", "kg / m²", 90, FMT_NUM, "GridMats sales pitch.")
inp_row("S1_Cement", "Cement in GridMats S1", "kg / m²", 3.7, FMT_NUM1, "GridMats sales pitch (3.7 t per 1,000 m²).")
inp_row("BitumenEF", "Bitumen binder emission factor", "t CO2e / t", 0.55, FMT_NUM2, "GridMats Design Engine factor.")
inp_row("HotMixEF", "Hot-mix production emissions", "kg CO2e / t mix", 25, FMT_NUM, "GridMats Design Engine factor.")
inp_row("BT_Mix", "Bituminous layers BM 50 + BC 40 mm", "t mix / m²", 0.2115, FMT_NUM2, "Derived: 90 mm × 2.35 t/m³.")
inp_row("BT_Binder", "Bitumen binder in those layers", "kg / m²", 10.6, FMT_NUM1, "Derived: about 5% binder by weight.")
inp_row("OV_Mix", "40 mm BC renewal", "t mix / m²", 0.094, FMT_NUM2, "Derived: 40 mm × 2.35 t/m³.")
inp_row("OV_Binder", "Binder in a 40 mm BC renewal", "kg / m²", 4.7, FMT_NUM1, "Derived: about 5% binder by weight.")

sec("6. GridMats Solar Walkway (pre-commercial: the pilot measures these)")
inp_row("Site", "Walkway pilot site", "choice", "Navi Mumbai", None,
        "Navi Mumbai = Reliance Corporate Park (higher power cost, most visible). Jamnagar = Giga Complex. Both are shown on the Solar Walkway sheet.",
        dv_list="Jamnagar,Navi Mumbai")
inp_row("PowerDens", "Rated power of the walkway", "W / m²", 160, FMT_NUM,
        "GridMats Solar Walkway flier (6 Oct 2026): 40 W per module, one on each grid, 4 grids per m².")
inp_row("WkYieldJam", "Yield, flat walk-on modules, Jamnagar", "kWh / kWp / yr", 1300, FMT_NUM,
        "Assumption: about 20% below a tilted Gujarat rooftop (~1,650) for flat mounting, heat and foot-traffic soiling.", key=True)
inp_row("WkYieldMum", "Yield, flat walk-on modules, Navi Mumbai", "kWh / kWp / yr", 1150, FMT_NUM,
        "Assumption: Mumbai rooftop 1,420-1,520 kWh/kWp, derated about 20%.")
inp_row("WkYieldMult", "Yield scenario multiplier", "×", (0.9, 1.0, 1.1), FMT_NUM2, "Scenario range on the yield above.")
inp_row("TariffJam", "Grid power cost avoided, Jamnagar (HT, landed)", "₹ / kWh", 7.50, FMT_INR2,
        "GERC FY27: HT energy charge ₹4.30/kWh plus fuel surcharge and duty; landed about ₹7.5-8.5 (estimate).")
inp_row("TariffMum", "Grid power cost avoided, Navi Mumbai (HT commercial, landed)", "₹ / kWh", 12.00, FMT_INR2,
        "MERC FY27: HT-II commercial ₹13.72/kVAh; HT-I industrial ₹8.44/kVAh.")
inp_row("TariffEsc", "Power tariff escalation", "% / yr", 0.03, FMT_PCT, "Assumption.")
inp_row("WkDegr", "Walkway module degradation", "% / yr", 0.007, FMT_PCT, "Assumption (higher than ground-mount for walk-on use).")
inp_row("WkLife", "Walkway module life", "years", 20, FMT_NUM, "Assumption; the pilot validates durability.")
inp_row("SolarLayer", "Walkable PV layer (modules + MLP lumber with wiring)", "₹ / m²", 9500, FMT_INR,
        "GridMats Solar pilot quote, Oct 2026: walkable module about ₹9,500/m².", key=True)
inp_row("BOS", "Inverter, cabling, protection, metering", "₹ / kWp", 12500, FMT_INR, "Assumption.")
inp_row("WkOM", "Walkway cleaning and inspection", "₹ / m² / yr", 150, FMT_INR, "Assumption.")
inp_row("InvRepl", "Inverter replacement", "₹ / kWp", 4000, FMT_INR, "Assumption.")
inp_row("InvYear", "Inverter replacement year", "year", 12, FMT_NUM, "Assumption.")
inp_row("ADRate", "Accelerated depreciation on solar (WDV)", "%", 0.40, FMT_PCT, "Income-tax Act: 40% WDV for solar plant.")
inp_row("RooftopCost", "Benchmark: C&I rooftop solar", "₹ / kWp", 40000, FMT_INR, "C&I rooftop turnkey ₹36-44/Wp (2026).")
inp_row("CarportCost", "Benchmark: solar carport", "₹ / kWp", 50000, FMT_INR, "Commercial carports ₹40,000-60,000/kW (2026).")

sec("7. Pilot (Reliance)")
inp_row("PilotMainKm", "Pilot main road, GridMats S1", "km", 0.4, FMT_NUM2, "Proposed pilot scope.", key=True)
inp_row("PilotTrackKm", "Pilot O&M track, GridMats S2", "km", 0.6, FMT_NUM2, "Proposed pilot scope.", key=True)
inp_row("PilotWalk", "Pilot GridMats Solar Walkway", "m²", 50, FMT_NUM, "Proposed pilot scope.", key=True)
inp_row("ListHD", "List price OC-HD 40 at pilot size (5,000-25,000 m²)", "₹ / m²", 935, FMT_INR, "GridMats price ladder (25 Sep 2026).")
inp_row("OfferHD", "Pilot offer price OC-HD 40", "₹ / m²", 900, FMT_INR, "Proposed: annual-contract rate at pilot size.", key=True)
inp_row("TrialPrice", "OC-HD 40 price on the 2025 Kutch trial", "₹ / m²", 1165, FMT_INR, "Reliance PO, Sep 2025.")
inp_row("LayRate", "Grid laid per crew-day", "m² / day", 350, FMT_NUM, "GridMats systems manual: 300-400 m² per crew of four.")
inp_row("ManDay", "Supervisor man-day rate", "₹ / day", 15000, FMT_INR, "GridMats standard rate.")
inp_row("TravelValue", "Supervisor travel and stay for the full build", "₹", 75000, FMT_INR, "Estimate; normally paid by the client.")
inp_row("DesignValue", "Site design (IITPAVE on site CBR), BoQ, method statement, hold-point QC", "₹", 150000, FMT_INR, "Estimate of value; included.")
inp_row("MonitorValue", "IoT sensors + live dashboard + soiling study + 6/12-month reports", "₹", 400000, FMT_INR, "Estimate of value; offered free.")
inp_row("CoFund", "GridMats co-funding of the solar walkway (PV layer + inverter)", "%", 0.50, FMT_PCT, "Proposed.", key=True)
inp_row("PilotCredit", "Pilot grid value credited against the first order of 25,000 m² or more", "%", 0.10, FMT_PCT, "Proposed.", key=True)
inp_row("Rebate", "Year-end rebate above 50,000 m² on a rate contract", "%", 0.03, FMT_PCT, "GridMats annual rate contract terms.")
inp_row("ContainerM2", "Grid per 40-ft high-cube container", "m²", 1100, FMT_NUM, "GridMats master data (OC-HD 40).")
ws.freeze_panes = "A7"

# ================================================================ UNIT RATES
ws = ur
title(ws, "Unit rates: build-up of installed cost (₹ / m², ex-GST)",
      "GridMats sections follow the GridMats base tables for Kutch's low-CBR (2-5%) saline ground. Blue quantities can be edited.")
widths(ws, {"A": 50, "B": 14, "C": 12, "D": 14, "E": 14, "F": 60})


def buildup(r0, head, items, total_name, note):
    section(ws, r0, head, ncols=6)
    header(ws, r0 + 1, ["Item", "Quantity / m²", "Unit", "Rate (₹)", "₹ / m²", "Basis"])
    r = r0 + 2
    first = r
    for lab, qty, unit, rate, basis in items:
        put(ws, f"A{r}", lab, kind="label")
        put(ws, f"B{r}", qty, FMT_NUM2 if isinstance(qty, (int, float)) and qty < 10 else FMT_NUM,
            kind="link" if isinstance(qty, str) else "input")
        put(ws, f"C{r}", unit, kind="note")
        put(ws, f"D{r}", f"={rate}", FMT_INR, kind="link")
        put(ws, f"E{r}", f"=B{r}*D{r}", FMT_INR)
        put(ws, f"F{r}", basis, kind="note")
        r += 1
    put(ws, f"A{r}", "Installed cost", kind="label", bold=True)
    put(ws, f"E{r}", f"=SUM(E{first}:E{r-1})", FMT_INR, bold=True)
    ws[f"E{r}"].border = TOP2
    put(ws, f"F{r}", note, kind="note")
    dname(total_name, "Unit Rates", f"E{r}")
    return r


r = buildup(4, "GridMats S1: cement-treated gravel, main / spine roads and inverter yards (load class D, CBR 2-5%)", [
    ("Subgrade preparation", 1, "m²", "SubgradeRate", "Compact to 98% MDD"),
    ("GSB 200 mm", 0.200, "m³", "GSBRate", "Base table, load class D, CBR 2-5%"),
    ("WMM 150 mm", 0.150, "m³", "WMMRate", "Base table, load class D"),
    ("Bedding 25 mm M-sand / crusher dust", 0.025, "m³", "SandRate", ""),
    ("Geotextile 200 gsm", 1, "m²", "Geotex", "Separator on saline formation"),
    ("GridMats OC-HD 40", 1, "m²", "GridHD", "Park / annual-contract rate"),
    ("Freight to site", 1, "m²", "Freight", "At cost"),
    ("Grid laying labour", 1, "m²", "LayLab", "Built by teams GridMats trains"),
    ("U-nails", "=NailDensity", "nos", "NailPrice", "1 per 2-3 m²"),
    ("Cement-treated gravel infill", 1, "m²", "CTGInfill", "Opens to trucks at 7 days (7-day cubes 2 MPa)"),
    ("Edges and drains", 1, "m²", "S1Edges", ""),
], "S1_Capex", "Compare: GridMats indicative ₹1,400-1,900/m²; Design Engine 73 internal road ₹2,011/m².")

r = buildup(r + 2, "GridMats S2: permeable stone chip, O&M tracks (load class C, CBR 2-5%) - as built at Kutch in 2025", [
    ("Subgrade preparation", 1, "m²", "SubgradeRate", ""),
    ("GSB 100 mm", 0.100, "m³", "GSBRate", "Base table, load class C, CBR 2-5%"),
    ("WMM 150 mm", 0.150, "m³", "WMMRate", "Base table, load class C"),
    ("Bedding 20 mm", 0.020, "m³", "SandRate", ""),
    ("Geotextile 200 gsm", 1, "m²", "Geotex", ""),
    ("GridMats OC-HD 40", 1, "m²", "GridHD", "Park / annual-contract rate"),
    ("Freight to site", 1, "m²", "Freight", "At cost"),
    ("Grid laying labour", 1, "m²", "LayLab", ""),
    ("U-nails", "=NailDensity", "nos", "NailPrice", ""),
    ("8-12 mm stone-chip infill (about 37 L/m²)", 0.037, "m³", "ChipRate", "Opens the same day it is filled"),
    ("Edges", 1, "m²", "S2Edges", ""),
], "S2_Capex", "Compare: GridMats recorded S2 range ₹1,434-1,866/m².")
r += 1
put(ws, f"A{r}", "Variant: S2 with lighter OC-LW 40 (cars and LCVs only)", kind="label")
put(ws, f"E{r}", "=S2_Capex-GridHD+GridLW", FMT_INR)
put(ws, f"F{r}", "For tracks without trucks or water tankers. Not used in the base case.", kind="note")
dname("S2LW_Capex", "Unit Rates", f"E{r}")

r = buildup(r + 3, "GridMats Solar Walkway: base and PV layer (per m² of walkway)", [
    ("Subgrade preparation", 1, "m²", "SubgradeRate", ""),
    ("Open-graded GSB 100 mm", 0.100, "m³", "OGSBRate", "Drains through"),
    ("Bedding 20 mm", 0.020, "m³", "SandRate", ""),
    ("GridMats OC-LW 40 (500 × 500 mm, 4 per m²)", 1, "m²", "GridLWSmall", "Under-1,000 m² price"),
    ("Freight", 1, "m²", "Freight", ""),
    ("Grid laying labour", 1, "m²", "LayLab", ""),
    ("Edges", 1, "m²", "S2Edges", ""),
], "WkBase", "GridMats structural base of the walkway.")
r += 1
put(ws, f"A{r}", "Walkable PV layer (modules + MLP lumber carrying the wiring)", kind="label")
put(ws, f"E{r}", "=SolarLayer", FMT_INR, kind="link")
r += 1
put(ws, f"A{r}", "Inverter, cabling, protection, metering", kind="label")
put(ws, f"E{r}", "=BOS*PowerDens/1000", FMT_INR)
put(ws, f"F{r}", "₹/kWp × kWp per m²", kind="note")
r += 1
put(ws, f"A{r}", "Solar walkway, installed", kind="label", bold=True)
put(ws, f"E{r}", f"=WkBase+E{r-2}+E{r-1}", FMT_INR, bold=True)
ws[f"E{r}"].border = TOP2
dname("WkTotal", "Unit Rates", f"E{r}")
r += 1
put(ws, f"A{r}", "Cost of the solar part per kWp", kind="label")
put(ws, f"E{r}", "=(SolarLayer+BOS*PowerDens/1000)/(PowerDens/1000)", FMT_INR)
put(ws, f"F{r}", "Compare: C&I rooftop ₹40,000/kWp; solar carport ₹40,000-60,000/kWp.", kind="note")
dname("WkPerKWp", "Unit Rates", f"E{r}")

r += 3
section(ws, r, "Conventional surfaces (from Inputs)", ncols=6)
header(ws, r + 1, ["Surface", "", "", "", "₹ / m²", "Typical section"])
conv = [("WBM / gravel road", "WBM_Capex", "GSB 150 + WBM 150; regrade every monsoon"),
        ("Bituminous road, IRC:37", "BT_Capex", "GSB 225 + WBM 200 + BM 50 + BC 40"),
        ("Concrete road, IRC:58", "PQC_Capex", "PQC 220 + DLC 150 (+ GSB); 14-28 days to open"),
        ("Paver-block walkway", "Paver_Capex", "60-80 mm pavers on sand")]
rr = r + 2
for lab, nm, sect in conv:
    put(ws, f"A{rr}", lab, kind="label")
    put(ws, f"E{rr}", f"={nm}", FMT_INR, kind="link")
    put(ws, f"F{rr}", sect, kind="note")
    rr += 1

# ================================================================ ROAD LCC (per m²)
ws = lcc
title(ws, "Road life-cycle cost per m² over 25 years",
      "₹ per m², nominal, escalated at the cost-escalation rate and discounted at the discount rate. Year 0 = build.")
widths(ws, {"A": 44, "B": 12})
for t in YEARS:
    ws.column_dimensions[YC[t]].width = 9
ws.column_dimensions[PVCOL].width = 12
put(ws, "A5", "Year", kind="label", bold=True)
for t in YEARS:
    put(ws, f"{YC[t]}5", t, FMT_NUM, bold=True, align="center")
    ws[f"{YC[t]}5"].fill = PatternFill("solid", fgColor=GREEN_LIGHT)
put(ws, f"{PVCOL}5", "PV", bold=True, align="center")
ws[f"{PVCOL}5"].fill = PatternFill("solid", fgColor=GREEN_LIGHT)
put(ws, "A6", "Discount factor", kind="label")
put(ws, "A7", "Cost index", kind="label")
for t in YEARS:
    c = YC[t]
    put(ws, f"{c}6", f"=1/(1+DiscRate)^{c}$5", '0.000')
    put(ws, f"{c}7", f"=(1+CostEsc)^{c}$5", '0.000')

LCC_ROWS = {}


def lcc_block(r0, head, rows_spec, key):
    """rows_spec: list of (label, year0_formula, yearN_formula with {c})."""
    section(ws, r0, head, ncols=29)
    r = r0 + 1
    first = r
    for lab, f0, fn in rows_spec:
        put(ws, f"A{r}", lab, kind="label")
        for t in YEARS:
            c = YC[t]
            f = f0 if t == 0 else fn
            put(ws, f"{c}{r}", f.format(c=c) if isinstance(f, str) else f, FMT_NUM)
        put(ws, f"{PVCOL}{r}", f"=SUMPRODUCT({FIRST}{r}:{LAST}{r},{FIRST}$6:{LAST}$6)", FMT_NUM, bold=True)
        LCC_ROWS[(key, lab)] = r
        r += 1
    put(ws, f"A{r}", "Total", kind="label", bold=True)
    for t in YEARS:
        c = YC[t]
        put(ws, f"{c}{r}", f"=SUM({c}{first}:{c}{r-1})", FMT_NUM, bold=True)
    put(ws, f"{PVCOL}{r}", f"=SUMPRODUCT({FIRST}{r}:{LAST}{r},{FIRST}$6:{LAST}$6)", FMT_NUM, bold=True)
    LCC_ROWS[(key, "total")] = r
    return r + 2


per = lambda nm, cyc: f"=IF(AND(MOD({{c}}$5,{cyc})=0,{{c}}$5<25),{nm}*{{c}}$7,0)"
r = 9
r = lcc_block(r, "WBM / gravel road", [
    ("Build", "=WBM_Capex", 0),
    ("Routine regrading", 0, "=WBM_Routine*{c}$7"),
    ("Regravelling", 0, per("WBM_Regravel", "WBM_Cycle")),
    ("Dust suppression (construction years)", 0, "=IF({c}$5<=DustYears,DustCost*{c}$7,0)"),
], "WBM")
r = lcc_block(r, "Bituminous road, IRC:37", [
    ("Build", "=BT_Capex", 0),
    ("Routine maintenance", 0, "=BT_Routine*{c}$7"),
    ("BC renewal", 0, per("BT_Overlay", "BT_Cycle")),
], "BT")
r = lcc_block(r, "Concrete road, IRC:58", [
    ("Build", "=PQC_Capex", 0),
    ("Routine maintenance", 0, "=PQC_Routine*{c}$7"),
    ("Joint reseal and slab repairs", 0, per("PQC_Repair", "PQC_Cycle")),
], "PQC")
r = lcc_block(r, "GridMats S1: cement-treated gravel", [
    ("Build", "=S1_Capex", 0),
    ("Routine upkeep", 0, "=GM_RoutinePct*S1_Capex*{c}$7"),
    ("Chip-and-cement top-up", 0, per("S1_TopUp", "S1_Cycle")),
], "S1")
r = lcc_block(r, "GridMats S2: permeable stone chip", [
    ("Build", "=S2_Capex", 0),
    ("Routine upkeep", 0, "=GM_RoutinePct*S2_Capex*{c}$7"),
    ("Chip top-up", 0, per("S2_TopUp", "S2_Cycle")),
    ("Dust suppression (residual)", 0, "=IF({c}$5<=DustYears,DustCost*(1-DustCutS2)*{c}$7,0)"),
], "S2")

section(ws, r, "Summary: 25-year life-cycle cost (present value, ₹ / m²)", ncols=8)
header(ws, r + 1, ["Surface", "", "Build", "PV of upkeep", "25-yr LCC (PV)", "₹ / sq ft", "vs build", ""])
SUM_ROWS = {}
rr = r + 2
for lab, key in [("WBM / gravel", "WBM"), ("Bituminous, IRC:37", "BT"), ("Concrete, IRC:58", "PQC"),
                 ("GridMats S1", "S1"), ("GridMats S2", "S2")]:
    tr = LCC_ROWS[(key, "total")]
    put(ws, f"A{rr}", lab, kind="label", bold=key in ("S1", "S2"))
    put(ws, f"C{rr}", f"={FIRST}{tr}", FMT_INR)
    put(ws, f"D{rr}", f"={PVCOL}{tr}-C{rr}", FMT_INR)
    put(ws, f"E{rr}", f"={PVCOL}{tr}", FMT_INR, bold=True)
    put(ws, f"F{rr}", f"=E{rr}/10.7639", FMT_INR)
    put(ws, f"G{rr}", f"=E{rr}/C{rr}", FMT_X)
    dname(f"LCC_{key}", "Road LCC", f"E{rr}")
    dname(f"Upk_{key}", "Road LCC", f"D{rr}")
    SUM_ROWS[key] = rr
    rr += 1
rr += 1
section(ws, rr, "Like-for-like comparisons (₹ / m²)", ncols=8)
header(ws, rr + 1, ["Comparison", "", "Extra build cost", "Upkeep saved (PV)", "Net 25-yr saving", "Saving %", "", ""])
rr += 2
for lab, a, b in [("O&M tracks: GridMats S2 vs WBM", "S2", "WBM"),
                  ("Main roads: GridMats S1 vs bituminous", "S1", "BT"),
                  ("Main roads: GridMats S1 vs concrete", "S1", "PQC")]:
    put(ws, f"A{rr}", lab, kind="label")
    put(ws, f"C{rr}", f"=C{SUM_ROWS[a]}-C{SUM_ROWS[b]}", FMT_INR)
    put(ws, f"D{rr}", f"=D{SUM_ROWS[b]}-D{SUM_ROWS[a]}", FMT_INR)
    put(ws, f"E{rr}", f"=E{SUM_ROWS[b]}-E{SUM_ROWS[a]}", FMT_INR, bold=True)
    put(ws, f"F{rr}", f"=E{rr}/E{SUM_ROWS[b]}", FMT_PCT)
    dname(f"Save_{a}_{b}", "Road LCC", f"E{rr}")
    dname(f"SavePct_{a}_{b}", "Road LCC", f"F{rr}")
    rr += 1
ws.freeze_panes = "C6"

# ================================================================ SOLAR PARK 100 MW
ws = park
title(ws, "Solar park block: GridMats roads vs conventional roads",
      "Per solar block (default 100 MWp). ₹ lakh unless stated. Positive = benefit to Reliance. Scenario set on Inputs.")
widths(ws, {"A": 56, "B": 14})
for t in YEARS:
    ws.column_dimensions[YC[t]].width = 9
ws.column_dimensions[PVCOL].width = 12
section(ws, 4, "Quantities", ncols=4)
q = [("Block size (MWp)", "=BlockMWp", FMT_NUM, None),
     ("Road length (km)", "=BlockMWp/100*RoadKm", FMT_NUM1, "ParkKm"),
     ("Main roads, GridMats S1 (m²)", "=BlockMWp/100*RoadKm*MainShare*1000*MainWidth", FMT_NUM, "MainArea"),
     ("O&M tracks, GridMats S2 (m²)", "=BlockMWp/100*RoadKm*(1-MainShare)*1000*TrackWidth", FMT_NUM, "TrackArea"),
     ("Total GridMats area (m²)", "=MainArea+TrackArea", FMT_NUM, "ParkArea"),
     ("Main roads would otherwise be", "=MainBaseline", None, None)]
for i, (lab, f, fmt, nm) in enumerate(q):
    rr = 5 + i
    put(ws, f"A{rr}", lab, kind="label")
    put(ws, f"B{rr}", f, fmt, kind="link" if i in (0, 5) else "calc")
    if nm:
        dname(nm, "Solar Park 100MW", f"B{rr}")

R = {}  # row registry
rr = 13
put(ws, f"A{rr}", "Year", kind="label", bold=True)
for t in YEARS:
    put(ws, f"{YC[t]}{rr}", t, FMT_NUM, bold=True, align="center")
    ws[f"{YC[t]}{rr}"].fill = PatternFill("solid", fgColor=GREEN_LIGHT)
put(ws, f"{PVCOL}{rr}", "PV / total", bold=True, align="center")
ws[f"{PVCOL}{rr}"].fill = PatternFill("solid", fgColor=GREEN_LIGHT)
YR = rr


def prow(key, label, f0, fn, fmt=FMT_NUM1, pv="pv", bold=False, kind="calc"):
    global rr
    rr += 1
    put(ws, f"A{rr}", label, kind="label", bold=bold)
    for t in YEARS:
        c = YC[t]
        f = f0 if t == 0 else fn
        put(ws, f"{c}{rr}", f.format(c=c, r=rr) if isinstance(f, str) else f, fmt, bold=bold, kind=kind)
    if pv == "pv":
        put(ws, f"{PVCOL}{rr}", f"=SUMPRODUCT({FIRST}{rr}:{LAST}{rr},{FIRST}${R['DF']}:{LAST}${R['DF']})", fmt, bold=True)
    elif pv == "sum":
        put(ws, f"{PVCOL}{rr}", f"=SUM({FIRST}{rr}:{LAST}{rr})", fmt, bold=True)
    R[key] = rr
    return rr


R["DF"] = YR + 1
prow("DF", "Discount factor", "=1/(1+DiscRate)^{c}$13", "=1/(1+DiscRate)^{c}$13", '0.000', pv=None)
prow("CI", "Cost index", "=(1+CostEsc)^{c}$13", "=(1+CostEsc)^{c}$13", '0.000', pv=None)
prow("EI", "Energy index (degradation × value escalation)", 0, "=(1-Degr)^({c}$13-1)*(1+EnergyEsc)^({c}$13-1)", '0.000', pv=None)
prow("KI", "Carbon value index", "=(1+CarbonEsc)^{c}$13", "=(1+CarbonEsc)^{c}$13", '0.000', pv=None)
rr += 1
section(ws, rr, "Road costs (₹ lakh)", ncols=29)
lr = lambda key, lab: LCC_ROWS[(key, lab)]
prow("ConvMain", "Conventional main roads",
     f"=MainArea*IF(MainBaseline=\"PQC\",'Road LCC'!{{c}}{lr('PQC','total')},'Road LCC'!{{c}}{lr('BT','total')})/100000",
     f"=MainArea*IF(MainBaseline=\"PQC\",'Road LCC'!{{c}}{lr('PQC','total')},'Road LCC'!{{c}}{lr('BT','total')})/100000")
prow("ConvTrack", "Conventional O&M tracks (WBM)",
     f"=TrackArea*'Road LCC'!{{c}}{lr('WBM','total')}/100000", f"=TrackArea*'Road LCC'!{{c}}{lr('WBM','total')}/100000")
prow("GMMain", "GridMats S1 main roads",
     f"=MainArea*'Road LCC'!{{c}}{lr('S1','total')}/100000", f"=MainArea*'Road LCC'!{{c}}{lr('S1','total')}/100000")
prow("GMTrack", "GridMats S2 O&M tracks",
     f"=TrackArea*'Road LCC'!{{c}}{lr('S2','total')}/100000", f"=TrackArea*'Road LCC'!{{c}}{lr('S2','total')}/100000")
prow("RoadSave", "Road cost saving (conventional − GridMats)",
     "={c}" + str(R["ConvMain"]) + "+{c}" + str(R["ConvTrack"]) + "-{c}" + str(R["GMMain"]) + "-{c}" + str(R["GMTrack"]),
     "={c}" + str(R["ConvMain"]) + "+{c}" + str(R["ConvTrack"]) + "-{c}" + str(R["GMMain"]) + "-{c}" + str(R["GMTrack"]),
     bold=True)
rr += 1
section(ws, rr, "Energy, cleaning and carbon", ncols=29)
prow("MWh", "Extra energy from cleaner modules (MWh)", 0,
     "=BlockMWp*Yield*SoilingLoss*DustShare*DustCutS2*{c}" + str(R["EI"]), FMT_NUM, pv="sum")
prow("EnergyVal", "Value of extra energy (₹ lakh)", 0, "={c}" + str(R["MWh"]) + "*1000*EnergyValue/100000")
prow("Clean", "Module-cleaning cost avoided (₹ lakh)", 0, "=BlockMWp*CleanCost*CleanSave*{c}" + str(R["CI"]))
prow("Early", "Earlier energisation (₹ lakh)", 0,
     "=IF({c}$13=1,DaysEarly*BlockMWp*Yield/365*1000*EnergyValue/100000,0)")
bt_ov = LCC_ROWS[("BT", "BC renewal")]
prow("CO2Mat", "CO2e avoided: plastic reuse and paving materials (t)",
     "=(ParkArea*PlasticHD*PPCredit+MainArea*IF(MainBaseline=\"PQC\",(PQC_Cement-S1_Cement)*CementEF,"
     "BT_Binder*BitumenEF+BT_Mix*HotMixEF-S1_Cement*CementEF))/1000",
     f"=IF(AND(MainBaseline<>\"PQC\",'Road LCC'!{{c}}{bt_ov}>0),MainArea*(OV_Binder*BitumenEF+OV_Mix*HotMixEF)/1000,0)",
     FMT_NUM, pv="sum")
prow("CO2En", "CO2e avoided: extra renewable energy (t)", 0, "={c}" + str(R["MWh"]) + "*GridEF", FMT_NUM, pv="sum")
prow("CarbMat", "Carbon value: materials (₹ lakh)",
     "={c}" + str(R["CO2Mat"]) + "*CarbonPrice*{c}" + str(R["KI"]) + "/100000",
     "={c}" + str(R["CO2Mat"]) + "*CarbonPrice*{c}" + str(R["KI"]) + "/100000")
prow("CarbEn", "Carbon value: extra energy (₹ lakh)", 0,
     "={c}" + str(R["CO2En"]) + "*CarbonPrice*{c}" + str(R["KI"]) + "/100000")
rr += 1
nb = "={c}%d+{c}%d+{c}%d+{c}%d+{c}%d+{c}%d" % (R["RoadSave"], R["EnergyVal"], R["Clean"], R["Early"], R["CarbMat"], R["CarbEn"])
prow("Net", "NET BENEFIT TO RELIANCE (₹ lakh)", nb, nb, bold=True)
prow("Cum", "Cumulative (₹ lakh)", "={c}%d" % R["Net"], "=OFFSET({c}%d,0,-1)+{c}%d" % (R["Net"] + 1, R["Net"]), pv=None)
# replace OFFSET with direct reference to previous column
for t in YEARS[1:]:
    c, p = YC[t], YC[t - 1]
    ws[f"{c}{R['Cum']}"] = f"={p}{R['Cum']}+{c}{R['Net']}"

net, cum, df, ci, ei, ki = (R[k] for k in ("Net", "Cum", "DF", "CI", "EI", "KI"))
rr += 2
section(ws, rr, "Results per block", ncols=4)
res = [
    ("NPV over 25 years (₹ Cr)", f"={PVCOL}{net}/100", FMT_CR, "Park_NPV"),
    ("NPV per MW (₹ lakh)", f"={PVCOL}{net}/BlockMWp", FMT_LAKH, "Park_NPVperMW"),
    ("Up-front GridMats premium (₹ Cr)", f"=-{FIRST}{R['RoadSave']}/100", FMT_CR, "Park_Premium"),
    ("Premium as % of solar plant capex", f"=Park_Premium/(BlockMWp*SolarCapex)", FMT_PCT, "Park_PremPct"),
    ("IRR on the premium", f"=IFERROR(IRR({FIRST}{net}:{LAST}{net}),\"no premium\")", FMT_PCT, "Park_IRR"),
    ("Simple payback (years)",
     f"=IF({FIRST}{cum}>=0,0,IF(COUNTIF({Y1}{cum}:{LAST}{cum},\"<0\")>=25,\"> 25\","
     f"COUNTIF({Y1}{cum}:{LAST}{cum},\"<0\")+(-INDEX({FIRST}{cum}:{LAST}{cum},COUNTIF({Y1}{cum}:{LAST}{cum},\"<0\")+1))"
     f"/INDEX({FIRST}{net}:{LAST}{net},COUNTIF({Y1}{cum}:{LAST}{cum},\"<0\")+2)))", FMT_YRS, "Park_Payback"),
    ("PV of road upkeep saved, years 1-25 (₹ Cr)", f"=({PVCOL}{R['RoadSave']}-{FIRST}{R['RoadSave']})/100", FMT_CR, "PV_Upkeep"),
    ("PV of extra energy (₹ Cr)", f"={PVCOL}{R['EnergyVal']}/100", FMT_CR, "PV_Energy"),
    ("PV of cleaning saved (₹ Cr)", f"={PVCOL}{R['Clean']}/100", FMT_CR, "PV_Clean"),
    ("PV of earlier energisation (₹ Cr)", f"={PVCOL}{R['Early']}/100", FMT_CR, "PV_Early"),
    ("PV of carbon value (₹ Cr)", f"=({PVCOL}{R['CarbMat']}+{PVCOL}{R['CarbEn']})/100", FMT_CR, "PV_Carbon"),
    ("Benefit-to-premium ratio", "=IFERROR((PV_Upkeep+PV_Energy+PV_Clean+PV_Early+PV_Carbon)/Park_Premium,\"n/a\")", FMT_X, "Park_BCR"),
    ("NPV at 8% discount rate (₹ Cr)", f"=(NPV(0.08,{Y1}{net}:{LAST}{net})+{FIRST}{net})/100", FMT_CR, "Park_NPV8"),
    ("NPV at 12% discount rate (₹ Cr)", f"=(NPV(0.12,{Y1}{net}:{LAST}{net})+{FIRST}{net})/100", FMT_CR, "Park_NPV12"),
    ("Extra clean energy over 25 years (MWh)", f"={PVCOL}{R['MWh']}", FMT_NUM, "Park_MWh"),
    ("CO2e avoided over 25 years (t)", f"={PVCOL}{R['CO2Mat']}+{PVCOL}{R['CO2En']}", FMT_NUM, "Park_CO2"),
    ("Recycled PP used (t)", "=ParkArea*PlasticHD/1000", FMT_NUM1, "Park_PP"),
    ("Road upkeep saved, years 1-25, undiscounted (₹ Cr)", f"=(SUM({Y1}{R['RoadSave']}:{LAST}{R['RoadSave']}))/100", FMT_CR, "Park_UpkeepNom"),
]
for i, (lab, f, fmt, nm) in enumerate(res):
    r2 = rr + 1 + i
    put(ws, f"A{r2}", lab, kind="label", bold=i in (0, 4, 5))
    put(ws, f"B{r2}", f, fmt, bold=i in (0, 4, 5))
    dname(nm, "Solar Park 100MW", f"B{r2}")
rr = rr + 1 + len(res) + 1
section(ws, rr, "Scenario building blocks (used by the Sensitivity sheet)", ncols=4)
units = [
    ("PV of road costs incl. build, per block (₹ Cr)", f"={PVCOL}{R['RoadSave']}/100", "U_Road"),
    ("PV of energy per unit dust share per ₹/kWh (₹ Cr)",
     f"=BlockMWp*Yield*SoilingLoss*DustCutS2*1000*SUMPRODUCT({Y1}{ei}:{LAST}{ei},{Y1}{df}:{LAST}{df})/10000000", "U_Energy"),
    ("PV of cleaning per 100% saving (₹ Cr)", f"=BlockMWp*CleanCost*SUMPRODUCT({Y1}{ci}:{LAST}{ci},{Y1}{df}:{LAST}{df})/100", "U_Clean"),
    ("PV of earlier energisation per day per ₹/kWh (₹ Cr)", f"=BlockMWp*Yield/365*1000*{Y1}{df}/10000000", "U_Early"),
    ("PV of materials carbon per ₹/t (₹ Cr)",
     f"=SUMPRODUCT({FIRST}{R['CO2Mat']}:{LAST}{R['CO2Mat']},{FIRST}{ki}:{LAST}{ki},{FIRST}{df}:{LAST}{df})/10000000", "U_CarbMat"),
    ("PV of energy carbon per unit dust share per ₹/t (₹ Cr)",
     f"=BlockMWp*Yield*SoilingLoss*DustCutS2*GridEF*SUMPRODUCT({Y1}{ei}:{LAST}{ei},{Y1}{ki}:{LAST}{ki},{Y1}{df}:{LAST}{df})/10000000", "U_CarbEn"),
    ("PV factor of escalating costs, years 1-25", f"=SUMPRODUCT({Y1}{ci}:{LAST}{ci},{Y1}{df}:{LAST}{df})", "U_PVF"),
    ("Check: rebuilt NPV − model NPV (should be 0)",
     "=U_Road+U_Energy*DustShare*EnergyValue+U_Clean*CleanSave+U_Early*DaysEarly*EnergyValue+U_CarbMat*CarbonPrice"
     "+U_CarbEn*DustShare*CarbonPrice-Park_NPV", "U_Check"),
]
for i, (lab, f, nm) in enumerate(units):
    r2 = rr + 1 + i
    put(ws, f"A{r2}", lab, kind="label")
    put(ws, f"B{r2}", f, '#,##0.0000')
    dname(nm, "Solar Park 100MW", f"B{r2}")
ws.freeze_panes = "C14"

# ================================================================ SOLAR WALKWAY
ws = walk
title(ws, "GridMats Solar Walkway: economics per m² of walkway",
      "Pre-commercial product. Rated power and price are from GridMats; yield and life are assumptions that the pilot will measure.")
widths(ws, {"A": 52, "B": 14})
for t in YEARS:
    ws.column_dimensions[YC[t]].width = 9
ws.column_dimensions[PVCOL].width = 12
section(ws, 4, "Inputs in use", ncols=4)
wq = [
    ("Site", "=Site", None, None),
    ("Rated power (kWp per m²)", "=PowerDens/1000", '0.000', "WkKWp"),
    ("Yield (kWh / kWp / yr)", "=IF(Site=\"Navi Mumbai\",WkYieldMum,WkYieldJam)*WkYieldMult", FMT_NUM, "WkYield"),
    ("Grid power cost avoided (₹ / kWh, year 1)", "=IF(Site=\"Navi Mumbai\",TariffMum,TariffJam)", FMT_INR2, "WkTariff"),
    ("Solar walkway installed (₹ / m²)", "=WkTotal", FMT_INR, None),
    ("Less: paver walkway it replaces (₹ / m²)", "=-Paver_Capex", FMT_INR, None),
    ("Extra investment for a generating walkway (₹ / m²)", "=WkTotal-Paver_Capex", FMT_INR, "WkInc"),
    ("Depreciable solar part (₹ / m²)", "=SolarLayer+BOS*WkKWp", FMT_INR, "WkSolarPart"),
]
for i, (lab, f, fmt, nm) in enumerate(wq):
    r2 = 5 + i
    put(ws, f"A{r2}", lab, kind="label", bold=i == 6)
    put(ws, f"B{r2}", f, fmt, kind="link" if i in (0, 4) else "calc", bold=i == 6)
    if nm:
        dname(nm, "Solar Walkway", f"B{r2}")
rr = 15
put(ws, f"A{rr}", "Year", kind="label", bold=True)
for t in YEARS:
    put(ws, f"{YC[t]}{rr}", t, FMT_NUM, bold=True, align="center")
    ws[f"{YC[t]}{rr}"].fill = PatternFill("solid", fgColor=GREEN_LIGHT)
put(ws, f"{PVCOL}{rr}", "PV / total", bold=True, align="center")
ws[f"{PVCOL}{rr}"].fill = PatternFill("solid", fgColor=GREEN_LIGHT)
W = {}


def wrow(key, label, f0, fn, fmt=FMT_NUM, pv="pv", bold=False):
    global rr
    rr += 1
    put(ws, f"A{rr}", label, kind="label", bold=bold)
    for t in YEARS:
        c = YC[t]
        f = f0 if t == 0 else fn
        put(ws, f"{c}{rr}", f.format(c=c) if isinstance(f, str) else f, fmt, bold=bold)
    if pv == "pv":
        put(ws, f"{PVCOL}{rr}", f"=SUMPRODUCT({FIRST}{rr}:{LAST}{rr},{FIRST}$16:{LAST}$16)", fmt, bold=True)
    elif pv == "sum":
        put(ws, f"{PVCOL}{rr}", f"=SUM({FIRST}{rr}:{LAST}{rr})", fmt, bold=True)
    W[key] = rr


wrow("DF", "Discount factor", "=1/(1+DiscRate)^{c}$15", "=1/(1+DiscRate)^{c}$15", '0.000', pv=None)
wrow("CI", "Cost index", "=(1+CostEsc)^{c}$15", "=(1+CostEsc)^{c}$15", '0.000', pv=None)
wrow("Gen", "Generation (kWh per m²)", 0, "=IF({c}$15<=WkLife,WkKWp*WkYield*(1-WkDegr)^({c}$15-1),0)", FMT_NUM1)
wrow("Tar", "Power cost avoided (₹ / kWh)", 0, "=WkTariff*(1+TariffEsc)^({c}$15-1)", FMT_INR2, pv=None)
wrow("Save", "Electricity bill saved (₹)", 0, "={c}%d*{c}%d" % (W["Gen"], W["Tar"]))
wrow("OM", "Cleaning and inspection (₹)", 0, "=IF({c}$15<=WkLife,-WkOM*{c}%d,0)" % W["CI"])
wrow("Inv", "Inverter replacement (₹)", 0, "=IF({c}$15=InvYear,-InvRepl*WkKWp*{c}%d,0)" % W["CI"])
wrow("Capex", "Extra investment (₹)", "=-WkInc", 0)
nf = "={c}%d+{c}%d+{c}%d+{c}%d" % (W["Save"], W["OM"], W["Inv"], W["Capex"])
wrow("Net", "Net cash flow, pre-tax (₹)", nf, nf, bold=True)
wrow("Cum", "Cumulative (₹)", "={c}%d" % W["Net"], "={c}%d" % W["Net"], pv=None)
for t in YEARS[1:]:
    ws[f"{YC[t]}{W['Cum']}"] = f"={YC[t-1]}{W['Cum']}+{YC[t]}{W['Net']}"
wrow("Open", "Tax written-down value, opening (₹)", "=WkSolarPart", "=WkSolarPart", pv=None)
for t in YEARS[1:]:
    ws[f"{YC[t]}{W['Open']}"] = f"={YC[t-1]}{W['Open']}-{YC[t-1]}{W['Open'] + 1}"
wrow("Dep", "Accelerated depreciation (₹)", "=0", "={c}%d*ADRate" % W["Open"], pv=None)
ws[f"{FIRST}{W['Dep']}"] = 0
# depreciation claimed from year 1 on the opening value
for t in YEARS[1:]:
    ws[f"{YC[t]}{W['Open']}"] = "=WkSolarPart" if t == 1 else f"={YC[t-1]}{W['Open']}-{YC[t-1]}{W['Dep']}"
wrow("AT", "Net cash flow, after tax with accelerated depreciation (₹)",
     "={c}%d" % W["Capex"],
     "=({c}%d+{c}%d+{c}%d)*(1-TaxRate)+{c}%d*TaxRate" % (W["Save"], W["OM"], W["Inv"], W["Dep"]), bold=True)

net, cum = W["Net"], W["Cum"]
rr += 2
section(ws, rr, "Results per m² of walkway", ncols=4)
wres = [
    ("Year-1 electricity bill saved (₹ / m²)", f"={Y1}{W['Save']}", FMT_INR, "Wk_Save1"),
    ("Year-1 generation (kWh / m²)", f"={Y1}{W['Gen']}", FMT_NUM1, "Wk_Gen1"),
    ("NPV, pre-tax (₹ / m²)", f"={PVCOL}{net}", FMT_INR, "Wk_NPV"),
    ("IRR, pre-tax", f"=IFERROR(IRR({FIRST}{net}:{LAST}{net}),\"n/a\")", FMT_PCT, "Wk_IRR"),
    ("IRR, after tax with 40% accelerated depreciation", f"=IFERROR(IRR({FIRST}{W['AT']}:{LAST}{W['AT']}),\"n/a\")", FMT_PCT, "Wk_IRRAT"),
    ("Simple payback (years)",
     f"=IF(COUNTIF({Y1}{cum}:{LAST}{cum},\"<0\")>=25,\"> 25\","
     f"COUNTIF({Y1}{cum}:{LAST}{cum},\"<0\")+(-INDEX({FIRST}{cum}:{LAST}{cum},COUNTIF({Y1}{cum}:{LAST}{cum},\"<0\")+1))"
     f"/INDEX({FIRST}{net}:{LAST}{net},COUNTIF({Y1}{cum}:{LAST}{cum},\"<0\")+2))", FMT_YRS, "Wk_Payback"),
    ("Levelised cost of the walkway's power (₹ / kWh)",
     f"=(WkSolarPart-{PVCOL}{W['OM']}-{PVCOL}{W['Inv']})/{PVCOL}{W['Gen']}", FMT_INR2, "Wk_LCOE"),
    ("Lifetime generation (kWh / m²)", f"=SUM({FIRST}{W['Gen']}:{LAST}{W['Gen']})", FMT_NUM, "Wk_GenLife"),
]
for i, (lab, f, fmt, nm) in enumerate(wres):
    r2 = rr + 1 + i
    put(ws, f"A{r2}", lab, kind="label", bold=i in (3, 4, 5))
    put(ws, f"B{r2}", f, fmt, bold=i in (3, 4, 5))
    dname(nm, "Solar Walkway", f"B{r2}")
rr = rr + 1 + len(wres) + 1
section(ws, rr, "Both sites side by side (simple payback on year-1 figures)", ncols=4)
header(ws, rr + 1, ["Site", "Year-1 saving net of O&M (₹ / m²)", "Simple payback (yrs)", ""])
for i, (lab, yv, tv) in enumerate([("Jamnagar", "WkYieldJam", "TariffJam"), ("Navi Mumbai", "WkYieldMum", "TariffMum")]):
    r2 = rr + 2 + i
    put(ws, f"A{r2}", lab, kind="label")
    put(ws, f"B{r2}", f"=WkKWp*{yv}*WkYieldMult*{tv}-WkOM", FMT_INR)
    put(ws, f"C{r2}", f"=WkInc/B{r2}", FMT_YRS)
    dname(f"Wk_SP_{'Jam' if i == 0 else 'Mum'}", "Solar Walkway", f"C{r2}")
    dname(f"Wk_S1_{'Jam' if i == 0 else 'Mum'}", "Solar Walkway", f"B{r2}")
rr += 5
section(ws, rr, "Cost per kWp compared", ncols=4)
for i, (lab, f) in enumerate([("GridMats Solar Walkway, solar part (₹ / kWp)", "=WkPerKWp"),
                              ("C&I rooftop solar (₹ / kWp)", "=RooftopCost"),
                              ("Solar carport (₹ / kWp)", "=CarportCost")]):
    put(ws, f"A{rr+1+i}", lab, kind="label")
    put(ws, f"B{rr+1+i}", f, FMT_INR, kind="link")
put(ws, f"A{rr+4}", "A walkway kWp costs more than a rooftop kWp. It earns its place where roofs are full or shaded and a walkway is needed anyway: it uses no extra land, roof or mounting.", kind="note")
rr += 6
section(ws, rr, "Pilot: GridMats Solar Walkway at the chosen site", ncols=4)
pw = [("Area (m²)", "=PilotWalk", FMT_NUM, "PW_Area"),
      ("Rated power (kWp)", "=PilotWalk*WkKWp", FMT_NUM1, "PW_kWp"),
      ("Year-1 generation (kWh)", "=PilotWalk*Wk_Gen1", FMT_NUM, "PW_kWh"),
      ("Year-1 electricity bill saved (₹)", "=PilotWalk*Wk_Save1", FMT_INR, "PW_Save"),
      ("Installed cost before GridMats co-funding (₹)", "=PilotWalk*WkTotal", FMT_INR, "PW_Cost")]
for i, (lab, f, fmt, nm) in enumerate(pw):
    put(ws, f"A{rr+1+i}", lab, kind="label")
    put(ws, f"B{rr+1+i}", f, fmt)
    dname(nm, "Solar Walkway", f"B{rr+1+i}")
ws.freeze_panes = "C16"

# ================================================================ SUSTAINABILITY
ws = sus
title(ws, "Sustainability impact", "Per solar block (default 100 MWp) over 25 years, and for the proposed pilot.")
widths(ws, {"A": 60, "B": 18, "C": 18, "D": 60})
header(ws, 4, ["Measure", "Per block", "Pilot", "Basis"])
pilot_main = "PilotMainKm*1000*MainWidth"
pilot_track = "PilotTrackKm*1000*TrackWidth"
srows = [
    ("Recycled polypropylene put to use (t)", "=Park_PP",
     f"=(({pilot_main})+({pilot_track}))*PlasticHD/1000+PilotWalk*PlasticLW/1000", "6.6 kg/m² OC-HD 40; 4.0 kg/m² OC-LW 40", "Sus_PP", "Sus_PP_P"),
    ("CO2e avoided: plastic reuse and paving materials (t)", f"='Solar Park 100MW'!{PVCOL}{R['CO2Mat']}",
     f"=(({pilot_main})+({pilot_track}))*PlasticHD*PPCredit/1000+({pilot_main})*IF(MainBaseline=\"PQC\",(PQC_Cement-S1_Cement)*CementEF,BT_Binder*BitumenEF+BT_Mix*HotMixEF-S1_Cement*CementEF)/1000",
     "Recycled-PP credit 1.2 kg/kg; bitumen or cement avoided on main roads, incl. BC renewals (block)", "Sus_CO2Mat", "Sus_CO2Mat_P"),
    ("CO2e avoided: extra renewable energy (t)", f"='Solar Park 100MW'!{PVCOL}{R['CO2En']}",
     "=PW_kWh*WkLife*GridEF/1000", "Block: cleaner modules × grid factor. Pilot: walkway generation over its life", "Sus_CO2En", "Sus_CO2En_P"),
    ("Total CO2e avoided (t)", "=B6+B7", "=C6+C7", "", "Sus_CO2", "Sus_CO2_P"),
    ("Extra clean energy (MWh)", "=Park_MWh", "=PW_kWh*WkLife/1000", "Block: 25 years. Pilot walkway: over its life", "Sus_MWh", "Sus_MWh_P"),
    ("Dust-suppression water avoided (kL)", "=TrackArea*DustWater*DustDays*DustYears*DustCutS2/1000",
     f"=({pilot_track})*DustWater*DustDays*DustYears*DustCutS2/1000", "Construction years only", "Sus_Water", "Sus_Water_P"),
    ("Rain returned to the ground every year (m³)", "=TrackArea*Rainfall/1000*Infiltr",
     f"=({pilot_track})*Rainfall/1000*Infiltr", "Permeable S2 tracks", "Sus_Rain", "Sus_Rain_P"),
    ("Cement avoided if main roads would otherwise be concrete (t)", "=MainArea*(PQC_Cement-S1_Cement)/1000",
     f"=({pilot_main})*(PQC_Cement-S1_Cement)/1000", "PQC 220 + DLC 150 ≈ 90 kg cement/m²; S1 uses 3.7 kg/m²", "Sus_Cement", "Sus_Cement_P"),
]
for i, (lab, fb, fp, basis, nb_, np_) in enumerate(srows):
    r2 = 5 + i
    put(ws, f"A{r2}", lab, kind="label", bold=i == 3)
    put(ws, f"B{r2}", fb, FMT_NUM, bold=i == 3)
    put(ws, f"C{r2}", fp, FMT_NUM1, bold=i == 3)
    put(ws, f"D{r2}", basis, kind="note")
    dname(nb_, "Sustainability", f"B{r2}")
    dname(np_, "Sustainability", f"C{r2}")
put(ws, "A15", "Circularity: GridMats grids are 100% post-consumer recycled polypropylene, and can be lifted, reused or recycled at end of life. "
    "They work alongside Reliance's own ReRoute plastic-in-bitumen roads: ReRoute for bituminous highways, GridMats where a road must be permeable, fast or cement-free.",
    kind="note")
ws["A15"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A15:D17")

# ================================================================ PILOT OFFER
ws = pilot
title(ws, "Pilot offer to Reliance: scope, value and price",
      "Ex-GST. List = GridMats standard price and terms. Offer = proposed pilot price. Civil works are by Reliance's EPC at its own rates.")
widths(ws, {"A": 74, "B": 12, "C": 9, "D": 13, "E": 15, "F": 13, "G": 15, "H": 50})
header(ws, 4, ["Item", "Qty", "Unit", "List rate (₹)", "List value (₹)", "Offer rate (₹)", "Offer value (₹)", "Notes"])
PR = {}
items = [
    ("A. Solar-park road pilot: 1 km with control sections", None),
    ("GridMats OC-HD 40, main road S1 (0.4 km × 6 m)", (f"={pilot_main}", "m²", "=ListHD", "=OfferHD", "Annual-contract rate at pilot size", "HDMain")),
    ("GridMats OC-HD 40, O&M track S2 (0.6 km × 4.5 m)", (f"={pilot_track}", "m²", "=ListHD", "=OfferHD", "", "HDTrack")),
    ("Geotextile 200 gsm", (f"=B6+B7", "m²", "=Geotex", "=Geotex", "", "Geo")),
    ("U-nails", (f"=ROUNDUP((B6+B7)*NailDensity,0)", "nos", "=NailPrice", "=NailPrice", "", "Nails")),
    ("Freight to site", (f"=B6+B7", "m²", "=Freight", "=Freight", "At cost", "Frt")),
    ("Site design (IITPAVE on measured CBR), BoQ, method statement, hold-point QC", (1, "lot", "=DesignValue", 0, "Included", "Des")),
    ("On-site GridMats supervisor for the whole build", (f"=ROUNDUP((B6+B7)/LayRate,0)", "days", "=ManDay", 0, "Standard is up to 3 days", "Sup")),
    ("Supervisor travel and stay", (1, "lot", "=TravelValue", 0, "Normally paid by the client", "Trv")),
    ("B. GridMats Solar Walkway demonstrator", None),
    ("GridMats base, OC-LW 40 on open-graded GSB", ("=PilotWalk", "m²", "=WkBase", "=WkBase", "", "WBase")),
    ("Walkable PV layer + inverter, cabling, metering", ("=PilotWalk", "m²", "=SolarLayer+BOS*PowerDens/1000",
                                                       "=(SolarLayer+BOS*PowerDens/1000)*(1-CoFund)", "GridMats co-funds the PV part", "WPV")),
    ("C. Measurement", None),
    ("IoT sensors, live dashboard, soiling study, 6- and 12-month joint reports", (1, "lot", "=MonitorValue", 0, "Free", "Mon")),
]
r = 5
for lab, spec in items:
    if spec is None:
        put(ws, f"A{r}", lab, kind="label", bold=True)
        ws[f"A{r}"].fill = PatternFill("solid", fgColor="F6F6F4")
        r += 1
        continue
    qty, unit, lr_, or_, note, key = spec
    put(ws, f"A{r}", lab, kind="label")
    put(ws, f"B{r}", qty, FMT_NUM, kind="input" if isinstance(qty, int) else "calc")
    put(ws, f"C{r}", unit, kind="note")
    put(ws, f"D{r}", lr_, FMT_INR, kind="link")
    put(ws, f"E{r}", f"=B{r}*D{r}", FMT_INR)
    put(ws, f"F{r}", or_, FMT_INR, kind="link" if isinstance(or_, str) else "input")
    put(ws, f"G{r}", f"=B{r}*F{r}", FMT_INR)
    put(ws, f"H{r}", note, kind="note")
    PR[key] = r
    r += 1
put(ws, f"A{r}", "Total GridMats scope", kind="label", bold=True)
put(ws, f"E{r}", f"=SUM(E5:E{r-1})", FMT_INR, bold=True)
put(ws, f"G{r}", f"=SUM(G5:G{r-1})", FMT_INR, bold=True)
for c in "EG":
    ws[f"{c}{r}"].border = TOP2
dname("Pilot_List", "Pilot Offer", f"E{r}")
dname("Pilot_Offer", "Pilot Offer", f"G{r}")
r += 2
section(ws, r, "What it means for Reliance", ncols=8)
summ = [
    ("Reliance pays GridMats (ex-GST)", "=Pilot_Offer", FMT_INR, "Pilot_Pays"),
    ("GridMats contribution (list value − offer)", "=Pilot_List-Pilot_Offer", FMT_INR, "Pilot_Contrib"),
    ("Saving on list value", "=Pilot_Contrib/Pilot_List", FMT_PCT, "Pilot_ContribPct"),
    ("Grid price reduction vs the 2025 Kutch trial (₹1,165/m²)", "=1-OfferHD/TrialPrice", FMT_PCT, "Pilot_vsTrial"),
    ("Civil works by Reliance's EPC, indicative (base layers, infill, labour, edges)",
     f"=({pilot_main})*(S1_Capex-GridHD-Geotex-Freight-NailDensity*NailPrice)+({pilot_track})*(S2_Capex-GridHD-Geotex-Freight-NailDensity*NailPrice)+PilotWalk*(WkBase-GridLWSmall-Freight)",
     FMT_INR, "Pilot_Civil"),
    ("Total pilot budget for Reliance (GridMats scope + civil)", "=Pilot_Pays+Pilot_Civil", FMT_INR, "Pilot_Budget"),
    ("Pilot road area (m²)", f"=({pilot_main})+({pilot_track})", FMT_NUM, "Pilot_RoadArea"),
    ("Grid laid in (crew-days)", "=ROUNDUP(Pilot_RoadArea/LayRate,0)", FMT_NUM, "Pilot_Days"),
    ("Credit against the first order of 25,000 m² or more", f"=PilotCredit*(G{PR['HDMain']}+G{PR['HDTrack']})", FMT_INR, "Pilot_Credit"),
    ("Net pilot cost to Reliance after that credit (GridMats scope)", "=Pilot_Pays-Pilot_Credit", FMT_INR, "Pilot_NetAfterCredit"),
    ("Rate held for a 12-month park rate contract (₹ / m² OC-HD 40)", "=OfferHD", FMT_INR, "Pilot_Lock"),
    ("Net 25-yr saving of the pilot road itself vs conventional (₹, PV)",
     f"=({pilot_main})*IF(MainBaseline=\"PQC\",Save_S1_PQC,Save_S1_BT)+({pilot_track})*Save_S2_WBM", FMT_INR, "Pilot_RoadNPV"),
]
for i, (lab, f, fmt, nm) in enumerate(summ):
    r2 = r + 1 + i
    put(ws, f"A{r2}", lab, kind="label", bold=i in (0, 5))
    put(ws, f"G{r2}", f, fmt, bold=i in (0, 5))
    dname(nm, "Pilot Offer", f"G{r2}")
r = r + 1 + len(summ) + 1
section(ws, r, "Success measures (12 months, measured jointly)", ncols=8)
kpis = [("Rut depth", "≤ 10 mm (investigate base if > 20 mm)"),
        ("Exposed cell tops", "≤ 5% of area"),
        ("Drainage", "No standing water 1 hour after rain"),
        ("Cracks", "None crossing more than 3 cells"),
        ("Road dust", "Dust deposition next to GridMats vs WBM control section"),
        ("Module soiling", "Soiling ratio of rows 10-40 m from GridMats vs WBM control"),
        ("Upkeep", "Maintenance events and cost vs control sections"),
        ("Solar walkway", "Measured W/m², kWh per m² per day, module temperature, slip resistance, uptime")]
for i, (k, v) in enumerate(kpis):
    put(ws, f"A{r+1+i}", k, kind="label")
    put(ws, f"B{r+1+i}", v, kind="note")
ws.freeze_panes = "A5"

# ================================================================ SCALE-UP
ws = scale
title(ws, "Scale-up: what the pilot unlocks", "Base on Solar Park results. Kutch: >40 billion kWh a year implies about 18-20 GW of solar.")
widths(ws, {"A": 52, "B": 15, "C": 15, "D": 15, "E": 15, "F": 40})
header(ws, 4, ["Measure", "Block", "1 GW", "5 GW", "20 GW", "Note"])
put(ws, "A5", "Solar capacity (MW)", kind="label", bold=True)
for c, v in zip("BCDE", ["=BlockMWp", 1000, 5000, 20000]):
    put(ws, f"{c}5", v, FMT_NUM, kind="link" if c == "B" else "input", bold=True)
srs = [("Internal road (km)", "={c}$5/100*RoadKm", FMT_NUM),
       ("GridMats area (m²)", "={c}$5/BlockMWp*ParkArea", FMT_NUM),
       ("Up-front GridMats premium (₹ Cr)", "={c}$5/BlockMWp*Park_Premium", FMT_CR),
       ("NPV to Reliance over 25 years (₹ Cr)", "={c}$5/BlockMWp*Park_NPV", FMT_CR),
       ("Road upkeep avoided, undiscounted (₹ Cr)", "={c}$5/BlockMWp*Park_UpkeepNom", FMT_CR),
       ("Extra clean energy over 25 years (GWh)", "={c}$5/BlockMWp*Park_MWh/1000", FMT_NUM1),
       ("Recycled PP put to use (t)", "={c}$5/BlockMWp*Park_PP", FMT_NUM),
       ("CO2e avoided over 25 years (t)", "={c}$5/BlockMWp*Park_CO2", FMT_NUM),
       ("40-ft containers of grid", "=ROUNDUP({c}$5/BlockMWp*ParkArea/ContainerM2,0)", FMT_NUM)]
SC = {}
for i, (lab, f, fmt) in enumerate(srs):
    r2 = 6 + i
    put(ws, f"A{r2}", lab, kind="label", bold=i == 3)
    for c in "BCDE":
        put(ws, f"{c}{r2}", f.format(c=c), fmt, bold=i == 3)
    SC[i] = r2
dname("Scale1GW_NPV", "Scale-up", f"C{SC[3]}")
dname("Scale1GW_PP", "Scale-up", f"C{SC[6]}")
dname("Scale1GW_CO2", "Scale-up", f"C{SC[7]}")
dname("Scale1GW_Km", "Scale-up", f"C{SC[0]}")
dname("Scale20GW_NPV", "Scale-up", f"E{SC[3]}")
dname("Scale20GW_Km", "Scale-up", f"E{SC[0]}")
section(ws, 17, "Solar walkways across Reliance sites (illustrative, year 1)", ncols=6)
header(ws, 18, ["Case", "Area (m²)", "Rated MWp", "MWh / yr", "₹ Cr / yr saved", "Basis"])
cases = [("Jamnagar Giga Complex and data-centre campus walkways", 10000, "WkYieldJam", "TariffJam", "Jamnagar yield and tariff"),
         ("Reliance Corporate Park and Jio World Centre walkways, Mumbai", 5000, "WkYieldMum", "TariffMum", "Mumbai yield and tariff"),
         ("Jio-bp EV-charging hubs: 500 sites × 40 m²", 20000, "WkYieldJam", "TariffJam", "Jamnagar basis, for scale only")]
for i, (lab, a, yv, tv, note) in enumerate(cases):
    r2 = 19 + i
    put(ws, f"A{r2}", lab, kind="label")
    put(ws, f"B{r2}", a, FMT_NUM, kind="input")
    put(ws, f"C{r2}", f"=B{r2}*WkKWp/1000", FMT_NUM2)
    put(ws, f"D{r2}", f"=B{r2}*WkKWp*{yv}*WkYieldMult/1000", FMT_NUM)
    put(ws, f"E{r2}", f"=D{r2}*1000*{tv}/10000000", FMT_NUM2)
    put(ws, f"F{r2}", note, kind="note")
r2 = 19 + len(cases)
put(ws, f"A{r2}", "Total", kind="label", bold=True)
for c, fmt in zip("BCDE", [FMT_NUM, FMT_NUM2, FMT_NUM, FMT_NUM2]):
    put(ws, f"{c}{r2}", f"=SUM({c}19:{c}{r2-1})", fmt, bold=True)
dname("Campus_MWp", "Scale-up", "C19")
dname("Campus_MWh", "Scale-up", "D19")
dname("Campus_Cr", "Scale-up", "E19")
dname("WkScale_Area", "Scale-up", f"B{r2}")
dname("WkScale_MWp", "Scale-up", f"C{r2}")
dname("WkScale_MWh", "Scale-up", f"D{r2}")
dname("WkScale_Cr", "Scale-up", f"E{r2}")

# ================================================================ SENSITIVITY
ws = sens
title(ws, "Sensitivity: NPV per solar block (₹ Cr, 25 years)",
      "Each row changes one driver and keeps everything else at the selected scenario. Low and high values can be edited.")
widths(ws, {"A": 52, "B": 13, "C": 13, "D": 15, "E": 15, "F": 13, "G": 40})
put(ws, "A4", "NPV at selected scenario (₹ Cr)", kind="label", bold=True)
put(ws, "D4", "=Park_NPV", FMT_CR, bold=True)
header(ws, 6, ["Driver", "Low value", "High value", "NPV at low", "NPV at high", "Swing", "How it is computed"])
grid_adj = "GridHD*ParkArea*(1+GM_RoutinePct*U_PVF)/10000000"
gm_upk = "(MainArea*Upk_S1+TrackArea*Upk_S2)/10000000"
road_npv = "(U_Road+U_CarbMat*CarbonPrice)"
tor = [
    ("GridMats grid price (± 10%)", -0.10, 0.10, FMT_PCT,
     f"=D$4-B{{r}}*{grid_adj}", f"=D$4-C{{r}}*{grid_adj}", "Build cost and the 0.5%/yr upkeep linked to it"),
    ("WBM upkeep (− / + 50%)", -0.5, 0.5, FMT_PCT,
     "=D$4+B{r}*TrackArea*Upk_WBM/10000000", "=D$4+C{r}*TrackArea*Upk_WBM/10000000", "Scales PV of WBM upkeep avoided"),
    ("Bituminous upkeep (− / + 50%)", -0.5, 0.5, FMT_PCT,
     "=D$4+B{r}*MainArea*Upk_BT/10000000*(MainBaseline<>\"PQC\")", "=D$4+C{r}*MainArea*Upk_BT/10000000*(MainBaseline<>\"PQC\")", "Only when main roads would be bituminous"),
    ("GridMats upkeep (+ / − 50%)", 0.5, -0.5, FMT_PCT,
     f"=D$4-B{{r}}*{gm_upk}", f"=D$4-C{{r}}*{gm_upk}", "Scales PV of GridMats upkeep"),
    ("Share of soiling from track dust", 0.05, 0.30, FMT_PCT,
     "=D$4+(B{r}-DustShare)*(U_Energy*EnergyValue+U_CarbEn*CarbonPrice)", "=D$4+(C{r}-DustShare)*(U_Energy*EnergyValue+U_CarbEn*CarbonPrice)", "Linear in the energy and energy-carbon terms"),
    ("Value of extra energy (₹ / kWh)", 2.40, 3.50, FMT_INR2,
     "=D$4+(B{r}-EnergyValue)*(U_Energy*DustShare+U_Early*DaysEarly)", "=D$4+(C{r}-EnergyValue)*(U_Energy*DustShare+U_Early*DaysEarly)", "Linear"),
    ("Carbon value (₹ / t CO2e)", 0, 1500, FMT_INR,
     "=D$4+(B{r}-CarbonPrice)*(U_CarbMat+U_CarbEn*DustShare)", "=D$4+(C{r}-CarbonPrice)*(U_CarbMat+U_CarbEn*DustShare)", "Linear"),
    ("Road length per 100 MW (km)", 8, 12, FMT_NUM,
     f"=D$4+(B{{r}}/RoadKm-1)*{road_npv}", f"=D$4+(C{{r}}/RoadKm-1)*{road_npv}", "Road terms scale with length"),
    ("Discount rate", 0.08, 0.12, FMT_PCT, "=Park_NPV8", "=Park_NPV12", "Full recomputation (8% and 12% fixed)"),
]
for i, (lab, lo, hi, fmt, flo, fhi, how) in enumerate(tor):
    r2 = 7 + i
    put(ws, f"A{r2}", lab, kind="label")
    fixed = lab == "Discount rate"
    put(ws, f"B{r2}", lo, fmt, kind="calc" if fixed else "input")
    put(ws, f"C{r2}", hi, fmt, kind="calc" if fixed else "input")
    put(ws, f"D{r2}", flo.format(r=r2), FMT_CR)
    put(ws, f"E{r2}", fhi.format(r=r2), FMT_CR)
    put(ws, f"F{r2}", f"=ABS(E{r2}-D{r2})", FMT_CR)
    put(ws, f"G{r2}", how, kind="note")
r = 18
section(ws, r, "Scenarios side by side", ncols=7)
header(ws, r + 1, ["Measure", "Conservative", "Base", "Optimistic", "", "", ""])
scn_rows = [("Share of soiling from track dust", "DustShare", FMT_PCT),
            ("Value of extra energy (₹ / kWh)", "EnergyValue", FMT_INR2),
            ("Cleaning cost avoided", "CleanSave", FMT_PCT),
            ("Earlier energisation (days)", "DaysEarly", FMT_NUM),
            ("Carbon value (₹ / t)", "CarbonPrice", FMT_INR)]
for i, (lab, nm, fmt) in enumerate(scn_rows):
    r2 = r + 2 + i
    put(ws, f"A{r2}", lab, kind="label")
    for c, ic in zip("BCD", "DEF"):
        put(ws, f"{c}{r2}", "=" + N[nm].replace("$C$", f"${ic}$"), fmt, kind="link")
r2 = r + 2 + len(scn_rows)
put(ws, f"A{r2}", "NPV per block (₹ Cr)", kind="label", bold=True)
b = r + 2
for c in "BCD":
    put(ws, f"{c}{r2}",
        f"=U_Road+U_Energy*{c}{b}*{c}{b+1}+U_Clean*{c}{b+2}+U_Early*{c}{b+3}*{c}{b+1}+U_CarbMat*{c}{b+4}+U_CarbEn*{c}{b}*{c}{b+4}",
        FMT_CR, bold=True)
dname("Scn_Cons", "Sensitivity", f"B{r2}")
dname("Scn_Base", "Sensitivity", f"C{r2}")
dname("Scn_Opt", "Sensitivity", f"D{r2}")
put(ws, f"A{r2+1}", "NPV per block if the energy benefit is zero (roads only) (₹ Cr)", kind="label")
put(ws, f"C{r2+1}", "=U_Road+U_CarbMat*CarbonPrice", FMT_CR)
dname("Scn_RoadsOnly", "Sensitivity", f"C{r2+1}")
put(ws, f"A{r2+3}", "Note: the conservative case still includes avoided WBM and bitumen upkeep. With no energy benefit at all, the roads-only line shows the NPV from upkeep and carbon alone.", kind="note")

# ================================================================ SOURCES
ws = src
title(ws, "Sources", "Public sources accessed Sep-Oct 2026; GridMats documents current at 7 Oct 2026.")
widths(ws, {"A": 6, "B": 70, "C": 90})
header(ws, 4, ["#", "Source", "Used for"])
sources = [
    ("GridMats master data and price ladder (Rev 1.3, 25-26 Sep 2026)", "Grid prices, panel weights, recycled PP per m², approved product claims"),
    ("GridMats Sales Pitch 2026 (Rev 1.5)", "Indicative installed costs, conventional comparators, cement and CO2 per 1,000 m²"),
    ("GridMats Systems Manual and specification clauses (Rev 1.3, Oct 2026)", "Base tables by load class and CBR, upkeep thresholds, laying rate"),
    ("GridMats Design Engine 73: technical deck, commercial deck, BoQ and sample proposal", "Unit rates, lifecycle method, emission factors, cross-sections"),
    ("GridMats solar and wind parks flier (Oct 2026)", "8-12 km of road per 100 MW, per-km park rate, system choice by road type"),
    ("GridMats Solar Walkway flier (6 Oct 2026) and GridMats Solar pilot quote (Oct 2026)", "160 W/m², construction, pilot method, ₹9,500/m²"),
    ("Reliance purchase and work orders, Kutch Z2 block trial (Sep 2025) and GridMats IITPAVE design (Aug 2025)", "2025 price, quantities, cross-section on CBR 3-5% saline ground"),
    ("Bureau Veritas BVIPL:SPL:BL/2545/7/2021/A; IIT Hyderabad interim report (Oct 2025)", "407 t/m² lab test on the empty OC-HD 40 panel; 2.0× base modulus at CBR 5%"),
    ("Reliance AGM 2026 coverage (BusinessToday, 19 Jun 2026); RIL Annual Report 2025-26; PL Capital Q1 FY27 note", "Kutch 5.5 lakh acres, >40 billion kWh/yr, 55 MWp/day peak; 100 GW by 2030; Net Carbon Zero 2035"),
    ("CEA CO2 Baseline Database v21.0 (FY2024-25)", "Grid emission factor 0.710 t/MWh"),
    ("GERC tariff orders FY2026-27 (25 Mar 2026); MERC Case 75 of 2025", "Gujarat and Maharashtra HT tariffs"),
    ("SECI ISTS-XX auction (Oct 2025); NHPC / NTPC solar tenders", "Solar PPA levels ₹2.52-2.87/kWh"),
    ("NHPC Khavda project pages; Gujarat plant yield data", "Kutch specific yield"),
    ("CPWD DSR 2023 BoQ via Infralens; Tamil Nadu G.O. 2021-22; IRC:82-2015", "Bituminous, concrete and WBM costs and upkeep"),
    ("Renewable Watch (2023); pv-magazine India (Apr 2026); taypro (2025); Univ. of Johannesburg thesis (2025)", "Cleaning cost, soiling rates, road-dust soiling gradient"),
    ("ICAP / reclimatize.in (2026) on India CCTS", "Carbon value range"),
    ("Greentechlead (2026), Bridgeway Power (2026), studiomatrx (2026)", "Solar capex, rooftop and carport costs"),
]
for i, (s, u) in enumerate(sources):
    put(ws, f"A{5+i}", i + 1, FMT_NUM)
    put(ws, f"B{5+i}", s, kind="label")
    put(ws, f"C{5+i}", u, kind="note")
put(ws, f"B{5+len(sources)+1}", "Figures marked 'Assumption' on the Inputs sheet are GridMats estimates for this model. The pilot is designed to measure the uncertain ones (soiling, dust, walkway yield).", kind="note")

# ================================================================ COVER
ws = cover
ws.sheet_view.showGridLines = False
widths(ws, {"A": 3, "B": 58, "C": 22, "D": 22, "E": 40})
ws["B2"] = "GridMats Solar × Reliance"
ws["B2"].font = Font(name=FONT, size=22, bold=True, color="C22935")
ws["B3"] = "Cost-benefit model for solar-park roads and solar walkways, with a pilot offer"
ws["B3"].font = Font(name=FONT, size=12, color="3F3F3F")
ws["B4"] = "Prepared for Reliance Industries Ltd (Reliance New Energy) by GroundReality Enterprises Pvt Ltd (GridMats), October 2026"
ws["B4"].font = NOTE_FONT
put(ws, "B6", "Scenario shown", kind="label", bold=True)
put(ws, "C6", "=Scenario", kind="link", bold=True)
put(ws, "D6", "Change it on the Inputs sheet (cell C3)", kind="note")
section(ws, 8, "Headline results", ncols=4, col=2)
header(ws, 9, ["Measure", "Value", "", "Where"], col=2)
cov = [
    ("Solar-park roads, per 100 MW block", None, None, None),
    ("NPV to Reliance over 25 years (₹ Cr)", "=Park_NPV", FMT_CR, "Solar Park 100MW"),
    ("Up-front GridMats premium (₹ Cr)", "=Park_Premium", FMT_CR, "Solar Park 100MW"),
    ("Premium as % of solar plant capex", "=Park_PremPct", FMT_PCT, "Solar Park 100MW"),
    ("IRR on the premium", "=Park_IRR", FMT_PCT, "Solar Park 100MW"),
    ("Simple payback (years)", "=Park_Payback", FMT_YRS, "Solar Park 100MW"),
    ("NPV, conservative case (₹ Cr)", "=Scn_Cons", FMT_CR, "Sensitivity"),
    ("NPV, optimistic case (₹ Cr)", "=Scn_Opt", FMT_CR, "Sensitivity"),
    ("25-year cost per m²", None, None, None),
    ("O&M tracks: GridMats S2 saving vs WBM", "=SavePct_S2_WBM", FMT_PCT, "Road LCC"),
    ("Main roads: GridMats S1 saving vs bituminous", "=SavePct_S1_BT", FMT_PCT, "Road LCC"),
    ("Main roads: GridMats S1 saving vs concrete", "=SavePct_S1_PQC", FMT_PCT, "Road LCC"),
    ("GridMats Solar Walkway, per m²", None, None, None),
    ("Pilot site", "=Site", None, "Inputs"),
    ("Year-1 electricity bill saved (₹ / m²)", "=Wk_Save1", FMT_INR, "Solar Walkway"),
    ("Simple payback (years)", "=Wk_Payback", FMT_YRS, "Solar Walkway"),
    ("IRR after tax with accelerated depreciation", "=Wk_IRRAT", FMT_PCT, "Solar Walkway"),
    ("Pilot offer", None, None, None),
    ("Reliance pays GridMats (ex-GST)", "=Pilot_Pays", FMT_INR, "Pilot Offer"),
    ("GridMats contribution (value given)", "=Pilot_Contrib", FMT_INR, "Pilot Offer"),
    ("Grid price reduction vs the 2025 Kutch trial", "=Pilot_vsTrial", FMT_PCT, "Pilot Offer"),
    ("Sustainability per 100 MW block", None, None, None),
    ("Recycled PP put to use (t)", "=Sus_PP", FMT_NUM, "Sustainability"),
    ("CO2e avoided over 25 years (t)", "=Sus_CO2", FMT_NUM, "Sustainability"),
]
r = 10
for lab, f, fmt, where in cov:
    if f is None:
        put(ws, f"B{r}", lab, kind="label", bold=True)
        for c in "BCDE":
            ws[f"{c}{r}"].fill = PatternFill("solid", fgColor="F6F6F4")
    else:
        put(ws, f"B{r}", lab, kind="label")
        put(ws, f"C{r}", f, fmt, kind="link", bold=True)
        put(ws, f"E{r}", where, kind="note")
    r += 1
r += 1
section(ws, r, "How to read this model", ncols=4, col=2)
notes = [
    "Blue cells are inputs, yellow-filled blue cells are key assumptions, black cells are formulas, green cells link to other sheets.",
    "All money is in Indian rupees, excluding GST, at 2026 prices, escalated and discounted as set on the Inputs sheet.",
    "The road comparison uses one solar block (default 100 MWp) with 10 km of internal road: 30% main roads (GridMats S1 instead of bituminous or concrete) and 70% O&M tracks (GridMats S2 instead of WBM).",
    "GridMats costs more than WBM on day one. The model shows where that premium is earned back: avoided regrading and regravelling, less dust on modules, and carbon.",
    "The GridMats Solar Walkway is at pilot stage. Its yield, life and upkeep are assumptions; the pilot measures them.",
    "Sources are listed on the Sources sheet and next to every input.",
]
for i, n in enumerate(notes):
    put(ws, f"B{r+1+i}", f"{i+1}. {n}", kind="label")
    ws[f"B{r+1+i}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(f"B{r+1+i}:E{r+1+i}")
    ws.row_dimensions[r + 1 + i].height = 30

for w in wb.worksheets:
    w.sheet_view.zoomScale = 90
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
